from __future__ import annotations

import hashlib
import json
import shutil
import sys
import traceback
from collections import Counter, defaultdict
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import create_engine, insert, select, text

PROJECT_ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(PROJECT_ROOT / "backend"))

from config import load_settings
from domain.daily_sales_schema import vip_daily_sales_table_for_year
from domain.factory_channel_sales_summary_schema import FACTORY_CHANNEL_SALES_DAILY_SUMMARY_TABLE
from storage.daily_sales_repository import DailySalesRepository, _header
from storage.date_normalization import parse_date
from storage.factory_channel_sales_summary_repository import FactoryChannelSalesSummaryRepository
from storage.vip_repository import _xlsx_sheet_rows


TARGET_DATE = date(2026, 10, 23)
EXPECTED_SHA256 = "0727598b011520b6225416c4403616430ffb4b8b6ef72ef79132aba40cf7e8a3"
AUDIT_DIRECTORY = Path(__file__).resolve().parent
REPORT_PATH = AUDIT_DIRECTORY / "report.json"


def write_report(report):
    REPORT_PATH.write_text(json.dumps(report, ensure_ascii=False, indent=2, default=str) + "\n", encoding="utf-8")


def read_source(source):
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        sheet_name = workbook.sheetnames[0]
    finally:
        workbook.close()
    iterator = _xlsx_sheet_rows(source)
    _, header_values = next(iterator)
    headers = {index: _header(value) for index, value in header_values.items() if _header(value)}
    required_headers = {"日期", "商品ID", "SIZE_ID", "货号", "销售量", "销售额"}
    if not required_headers <= set(headers.values()):
        raise ValueError("Source headers do not match the VIP daily-sales format")
    rows = []
    keys = set()
    for row_number, values in iterator:
        if not any(values.values()):
            continue
        raw = {header: values.get(index, "") for index, header in headers.items()}
        sales_date = parse_date(raw.get("日期"))
        if sales_date != TARGET_DATE:
            raise ValueError(f"Unexpected sales date on row {row_number}: {sales_date}")
        row = DailySalesRepository._map_vip(raw)
        if not all(row[field] for field in ("goods_id", "size_id", "goods_code")):
            raise ValueError(f"Missing business identifiers on row {row_number}")
        if row["sales_quantity"] is None or row["sales_amount"] is None:
            raise ValueError(f"Invalid sales values on row {row_number}")
        for header in ("销售量", "客户数", "在售库存"):
            value = raw.get(header)
            if value and Decimal(value) != Decimal(value).to_integral_value():
                raise ValueError(f"Non-integer value on row {row_number}: {header}")
        row.update({
            "sales_date": sales_date,
            "source_workbook": source.name,
            "source_sheet": sheet_name,
            "source_row_number": row_number,
            "raw_payload": raw,
        })
        key = DailySalesRepository._vip_key(row)
        if key in keys:
            raise ValueError(f"Duplicate business key on row {row_number}: {key}")
        keys.add(key)
        rows.append(row)
    if not rows:
        raise ValueError("Source file contains no data")
    return rows


def source_totals(rows):
    brands = defaultdict(lambda: {"rows": 0, "sales_quantity": 0, "sales_amount": Decimal("0")})
    for row in rows:
        totals = brands[row["brand_name"] or ""]
        totals["rows"] += 1
        totals["sales_quantity"] += row["sales_quantity"] or 0
        totals["sales_amount"] += row["sales_amount"] or Decimal("0")
    return {
        "rows": len(rows),
        "unique_business_keys": len({DailySalesRepository._vip_key(row) for row in rows}),
        "goods_codes": len({row["goods_code"] for row in rows}),
        "sales_quantity": sum(row["sales_quantity"] or 0 for row in rows),
        "sales_amount": sum((row["sales_amount"] or Decimal("0") for row in rows), Decimal("0")),
        "customer_count": sum(row["customer_count"] or 0 for row in rows),
        "on_sale_stock": sum(row["on_sale_stock"] or 0 for row in rows),
        "brands": dict(brands),
    }


def unaffected_dates(connection):
    queries = {
        "vip_nearby": """
            SELECT sales_date, count(*) AS rows, sum(sales_quantity) AS sales_quantity,
                   sum(sales_amount) AS sales_amount
            FROM vip_daily_sales_2026
            WHERE sales_date BETWEEN DATE '2026-10-21' AND DATE '2026-10-24'
              AND sales_date <> :target_date
            GROUP BY sales_date ORDER BY sales_date
        """,
        "jst_target": """
            SELECT count(*) AS rows, sum(net_sales_quantity) AS net_sales_quantity,
                   sum(sales_amount) AS sales_amount, sum(net_sales_amount) AS net_sales_amount
            FROM jst_daily_sales_2026 WHERE sales_date = :target_date
        """,
        "summary_nearby": """
            SELECT sales_date, count(*) AS rows, sum(quantity) AS quantity,
                   sum(gross_quantity) AS gross_quantity, sum(return_quantity) AS return_quantity
            FROM factory_channel_sales_daily_summaries
            WHERE sales_date BETWEEN DATE '2026-10-21' AND DATE '2026-10-24'
              AND sales_date <> :target_date
            GROUP BY sales_date ORDER BY sales_date
        """,
    }
    return {
        name: [dict(row) for row in connection.execute(text(query), {"target_date": TARGET_DATE}).mappings()]
        for name, query in queries.items()
    }


def verify_rows(connection, table, rows):
    expected = {DailySalesRepository._vip_key(row): row for row in rows}
    stored = list(connection.execute(select(table).where(table.c.sales_date == TARGET_DATE)).mappings())
    actual_keys = {DailySalesRepository._vip_key(row) for row in stored}
    if actual_keys != set(expected) or len(stored) != len(expected):
        raise ValueError("Stored business keys do not exactly match the source")
    for actual in stored:
        original = expected[DailySalesRepository._vip_key(actual)]
        for field, value in original.items():
            if actual[field] != value:
                raise ValueError(f"Stored field differs from source: {actual['source_row_number']}, {field}")
    totals = source_totals(stored)
    if totals != source_totals(rows):
        raise ValueError("Stored totals differ from source totals")
    return {"mismatched_fields": 0, "totals": totals}


def main():
    report = {
        "status": "validating",
        "started_at": datetime.now().astimezone().isoformat(),
        "target_date": TARGET_DATE.isoformat(),
        "data_committed": False,
    }
    engine = None
    summary_repository = None
    try:
        sources = list((Path.home() / "Desktop").glob("*20261023-20261023_202610251009*.xlsx"))
        if len(sources) != 1:
            raise ValueError(f"Expected exactly one source workbook, found {len(sources)}")
        source = sources[0]
        digest = hashlib.sha256(source.read_bytes()).hexdigest()
        if digest != EXPECTED_SHA256:
            raise ValueError("Source workbook changed after preview")
        rows = read_source(source)
        totals = source_totals(rows)
        if totals["rows"] != 51710 or totals["sales_quantity"] != 2502 or totals["sales_amount"] != Decimal("590463.79"):
            raise ValueError("Source totals differ from the approved preview")
        report.update({"source": str(source), "source_sha256": digest, "source_totals": totals})
        write_report(report)
        if "--execute" not in sys.argv:
            print(json.dumps(report, ensure_ascii=False, default=str))
            return 0
        audit_source = AUDIT_DIRECTORY / source.name
        if not audit_source.exists():
            shutil.copy2(source, audit_source)
        if hashlib.sha256(audit_source.read_bytes()).hexdigest() != digest:
            raise ValueError("Audit copy does not match the source")
        settings = load_settings(require_database=True)
        engine = create_engine(
            settings.database_url,
            future=True,
            connect_args={"connect_timeout": 10, "application_name": "manual_vip_daily_backfill_20261023"},
        )
        table = vip_daily_sales_table_for_year(TARGET_DATE.year)
        with engine.begin() as connection:
            connection.execute(text("SET LOCAL lock_timeout = '5s'"))
            connection.execute(text("SET LOCAL statement_timeout = '120s'"))
            connection.execute(text("SELECT pg_advisory_xact_lock(hashtext('manual_vip_daily_backfill_20261023'))"))
            existing = connection.execute(
                text("SELECT count(*) FROM vip_daily_sales_2026 WHERE sales_date = :target_date"),
                {"target_date": TARGET_DATE},
            ).scalar_one()
            report["existing_rows_before"] = existing
            if existing:
                raise ValueError(f"Target date already has {existing} rows; refusing to overwrite data")
            report["unaffected_before"] = unaffected_dates(connection)
            report["task_status_before"] = [dict(row) for row in connection.execute(text(
                "SELECT * FROM scheduled_task_statuses WHERE task_name = 'import_vip_daily_sales' AND business_date = :target_date"
            ), {"target_date": TARGET_DATE}).mappings()]
            report["summary_before"] = [dict(row) for row in connection.execute(
                select(FACTORY_CHANNEL_SALES_DAILY_SUMMARY_TABLE)
                .where(FACTORY_CHANNEL_SALES_DAILY_SUMMARY_TABLE.c.sales_date == TARGET_DATE)
            ).mappings()]
            report["status"] = "inserting"
            write_report(report)
            for offset in range(0, len(rows), 1000):
                connection.execute(insert(table), rows[offset:offset + 1000])
            report["verification"] = verify_rows(connection, table, rows)
        report.update({"data_committed": True, "inserted_rows": len(rows), "status": "refreshing_summary"})
        write_report(report)
        print(f"Inserted and verified {len(rows)} VIP daily-sales rows", flush=True)
        summary_repository = FactoryChannelSalesSummaryRepository(settings.database_url)
        report["summary_refresh"] = summary_repository.refresh(
            sales_year=TARGET_DATE.year,
            date_start=TARGET_DATE,
            date_end=TARGET_DATE,
        )
        with engine.connect() as connection, connection.begin():
            connection.execute(text("SET TRANSACTION READ ONLY"))
            connection.execute(text("SET LOCAL statement_timeout = '120s'"))
            report["verification"] = verify_rows(connection, table, rows)
            report["unaffected_after"] = unaffected_dates(connection)
            if report["unaffected_after"] != report["unaffected_before"]:
                raise ValueError("Nearby dates or JST source totals changed during the backfill")
            report["summary_after"] = [dict(row) for row in connection.execute(text("""
                SELECT brand, match_status, count(*) AS rows, sum(quantity) AS quantity
                FROM factory_channel_sales_daily_summaries WHERE sales_date = :target_date
                GROUP BY brand, match_status ORDER BY brand, match_status
            """), {"target_date": TARGET_DATE}).mappings()]
            summary_count = sum(row["rows"] for row in report["summary_after"])
            if summary_count != report["summary_refresh"]["written"]:
                raise ValueError("Summary refresh row count does not match stored data")
        report.update({"status": "completed", "completed_at": datetime.now().astimezone().isoformat()})
        write_report(report)
        print(json.dumps({key: value for key, value in report.items() if key not in {"summary_before", "task_status_before"}}, ensure_ascii=False, default=str))
        print(f"Audit report: {REPORT_PATH}")
        return 0
    except Exception as error:
        report.update({"status": "failed", "error": f"{type(error).__name__}: {error}", "traceback": traceback.format_exc()})
        write_report(report)
        raise
    finally:
        if engine is not None:
            engine.dispose()
        if summary_repository is not None:
            summary_repository.engine.dispose()


if __name__ == "__main__":
    raise SystemExit(main())
