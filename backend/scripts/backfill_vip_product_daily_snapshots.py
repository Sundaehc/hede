from __future__ import annotations

import argparse
from datetime import date, timedelta
import json
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select

from config import load_settings
from domain.vip_schema import VIP_DAILY_SNAPSHOT_TABLE
from domain.vip_sources import VIP_DAILY_COLUMN_ALIASES
from storage.date_normalization import parse_date_range
from storage.vip_daily_snapshot_backfill import backfill_snapshot_file
from storage.vip_repository import VipRepository


def source_report_date(path: Path) -> date | None:
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        iterator = workbook.active.iter_rows(values_only=True)
        headers = next(iterator, ())
        date_index = next((index for index, header in enumerate(headers) if VIP_DAILY_COLUMN_ALIASES.get(str(header).strip()) == "date"), None)
        if date_index is None:
            return None
        for values in iterator:
            if date_index < len(values) and values[date_index] is not None:
                start_date, end_date = parse_date_range(values[date_index])
                return end_date if start_date == end_date else None
        return None
    finally:
        workbook.close()


def main() -> int:
    parser = argparse.ArgumentParser(description="Backfill snapshots only from reports with matching business dates")
    parser.add_argument("--dates", nargs="+", type=date.fromisoformat, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--execute", action="store_true")
    args = parser.parse_args()
    settings = load_settings(require_database=True)
    repository = VipRepository(settings.database_url)
    targets = set(args.dates)
    if not settings.vip_data_roots:
        raise ValueError("VIP source directories are not configured")
    report = {"execute": args.execute, "targets": sorted(day.isoformat() for day in targets), "files": [], "unresolved": []}
    args.output.parent.mkdir(parents=True, exist_ok=True)
    scan_days = sorted({day + timedelta(days=offset) for day in targets for offset in (0, 1, 2)})
    try:
        for root in settings.vip_data_roots:
            matched = set()
            for folder_date in scan_days:
                folder = root / folder_date.strftime("%m.%d")
                for extension in (".xlsx", ".xlsm"):
                    for path in sorted(folder.glob(f"*{extension}")):
                        if path.name.startswith("~$") or "日罗盘" not in path.stem:
                            continue
                        actual_date = source_report_date(path)
                        if actual_date not in targets:
                            report["files"].append({"source_file": str(path), "folder_date": folder_date.isoformat(), "report_date": actual_date.isoformat() if actual_date else None, "status": "date_mismatch"})
                            continue
                        if actual_date in matched:
                            continue
                        result = backfill_snapshot_file(repository, path, actual_date, execute=args.execute)
                        result["status"] = "inserted" if args.execute else "ready"
                        result["source_root"] = str(root)
                        report["files"].append(result)
                        if result["source_rows"]:
                            matched.add(actual_date)
                        print(json.dumps({key: value for key, value in result.items() if key != "inserted_ids"}, ensure_ascii=False), flush=True)
                        args.output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
            report["unresolved"].extend({"source_root": str(root), "snapshot_date": day.isoformat()} for day in sorted(targets - matched))
        with repository.engine.connect() as connection:
            present = set(connection.execute(
                select(VIP_DAILY_SNAPSHOT_TABLE.c.snapshot_date)
                .where(VIP_DAILY_SNAPSHOT_TABLE.c.snapshot_date.in_(targets)).distinct()
            ).scalars())
        report["database_missing_dates"] = sorted(day.isoformat() for day in targets - present)
    except Exception as exc:
        report["error"] = f"{type(exc).__name__}: {exc}"
        raise
    finally:
        args.output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        repository.engine.dispose()
    print(f"[OUTPUT] {args.output}; missing={report['database_missing_dates']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
