from __future__ import annotations

import io
from collections import OrderedDict
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP

import xlrd
from fastapi import HTTPException
from openpyxl import load_workbook

from domain.inventory_sources import ACCOUNTING_DOCUMENT_TYPES
from domain.inventory_numeric import INVENTORY_NUMBER_MAX, INVENTORY_NUMBER_OVERFLOW


TEMPLATE_HEADERS = {
    "purchase": ("日期", "单据类型", "单位全名", "仓库全名", "经手人", "摘要", "商品编码", "数量", "单价"),
    "purchase_return": ("日期", "单据类型", "单位全名", "仓库全名", "经手人", "摘要", "商品编码", "数量", "单价"),
    "sale": ("日期", "单据类型", "单位全名", "仓库全名", "经手人", "摘要", "商品编码", "数量", "单价"),
    "sale_return": ("日期", "单据类型", "单位全名", "仓库全名", "经手人", "摘要", "商品编码", "数量", "单价"),
    "stock_loss": ("日期", "单据类型", "仓库全名", "经手人", "摘要", "商品编码", "数量"),
    "stock_gain": ("日期", "单据类型", "仓库全名", "经手人", "摘要", "商品编码", "数量"),
    "transfer": ("日期", "单据类型", "出货仓库", "入货仓库", "经手人", "摘要", "商品编码", "数量"),
    "accounting": ("日期", "单据类型", "经手人", "单位全名", "摘要", "费用项目名", "总金额"),
}
LEGACY_TEMPLATE_HEADERS = {
    "purchase": ("单据日期", "单据类型", "单位全名", "仓库全名", "制单人", "摘要", "商品编码", "数量", "单价"),
    "purchase_return": ("日期", "单据类型", "单位全名", "仓库全名", "制单人", "摘要", "商品编码", "数量", "单价"),
    "sale": ("日期", "单据类型", "单位全名", "仓库全名", "经手人", "摘要", "商品编码", "销售数量", "单价"),
    "accounting": ("日期", "单据类型", "制单人", "往来单位全名", "摘要", "费用项目名", "总金额"),
}
LEGACY_NUMBER_HEADER = "单据编号"
TEMPLATE_TYPES = {
    "purchase": {"进货单"},
    "purchase_return": {"进货退货单"},
    "sale": {"批发销售单"},
    "sale_return": {"批发销售退货单"},
    "stock_loss": {"报损单"},
    "stock_gain": {"报溢单"},
    "transfer": {"同价调拨单"},
    "accounting": set(ACCOUNTING_DOCUMENT_TYPES),
}

def validate_template_number(value: object, label: str) -> Decimal:
    try:
        number = Decimal(str(value))
    except (InvalidOperation, ValueError):
        raise ValueError(f"{label}必须是有效数字") from None
    if not number.is_finite():
        raise ValueError(f"{label}必须是有效数字")
    if number.copy_abs() >= INVENTORY_NUMBER_OVERFLOW:
        raise ValueError(
            f"{label}超出系统数值范围（绝对值不能超过 {INVENTORY_NUMBER_MAX}），"
            "请核对数量、单价和金额；单据汇总过大时请使用不同摘要拆分单据"
        )
    return number.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def validate_template_document_numbers(rows: list[dict[str, object]], label: str) -> None:
    total_quantity = Decimal("0")
    total_amount = Decimal("0")
    for index, row in enumerate(rows, start=1):
        row_label = f"{label}：明细第 {index} 行"
        if row.get("product_code"):
            row_label += f"（商品编码：{row['product_code']}）"
        values = {}
        for field, field_label in (("quantity", "数量"), ("unit_price", "单价"), ("amount", "金额")):
            value = row.get(field)
            if value is not None and value != "":
                values[field] = validate_template_number(value, f"{row_label}的{field_label}")
        if "amount" not in row and "quantity" in values and "unit_price" in values:
            amount = Decimal(str(row["quantity"])) * Decimal(str(row["unit_price"]))
            values["amount"] = validate_template_number(amount, f"{row_label}的金额")
        total_quantity += values.get("quantity", Decimal("0"))
        total_amount += values.get("amount", Decimal("0"))
    validate_template_number(total_quantity, f"{label}的总数量")
    validate_template_number(total_amount, f"{label}的总金额")


@dataclass
class TemplateDocument:
    key: str
    date: str
    document_type: str
    supplier: str
    warehouse: str
    handler: str
    summary: str
    rows: list[dict[str, str]]


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _date(value: object, datemode: int) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    raw = _text(value)
    try:
        if isinstance(value, (int, float)) or (raw.replace(".", "", 1).isdigit() and float(raw) > 1000):
            return xlrd.xldate.xldate_as_datetime(float(raw), datemode).date().isoformat()
        return date.fromisoformat(raw.replace("/", "-").replace(".", "-").split()[0]).isoformat()
    except (ValueError, TypeError, OverflowError, IndexError):
        raise HTTPException(status_code=400, detail=f"无效日期：{raw}") from None


def _amount(value: object, row_number: int, field: str, *, required: bool = True) -> str:
    raw = _text(value)
    if not raw and not required:
        return ""
    try:
        number = Decimal(raw)
        if not number.is_finite() or number < 0 or (required and number == 0):
            raise InvalidOperation
    except (InvalidOperation, ValueError):
        raise HTTPException(status_code=400, detail=f"Excel 第 {row_number} 行：{field}必须是有效的正数") from None
    try:
        validate_template_number(number, f"Excel 第 {row_number} 行的{field}")
    except ValueError as error:
        raise HTTPException(status_code=400, detail=str(error)) from error
    return str(number)


def read_template_documents(content: bytes) -> tuple[str, str, list[TemplateDocument]]:
    workbook = None
    try:
        if content.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
            workbook = xlrd.open_workbook(file_contents=content)
            worksheet = workbook.sheet_by_index(0)
            rows = (worksheet.row_values(index) for index in range(worksheet.nrows))
            sheet_name = worksheet.name
            datemode = workbook.datemode
        else:
            workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
            worksheet = workbook.active
            rows = worksheet.iter_rows(values_only=True)
            sheet_name = worksheet.title
            datemode = 1 if workbook.epoch.year == 1904 else 0
        headers = tuple(_text(cell) for cell in next(rows, ()))
        import_headers = headers[:-1] if headers and headers[-1] == LEGACY_NUMBER_HEADER else headers
        matched_kinds = {
            key
            for template_headers in (TEMPLATE_HEADERS, LEGACY_TEMPLATE_HEADERS)
            for key, columns in template_headers.items()
            if import_headers == columns or (
                key in {"sale", "sale_return"}
                and import_headers == tuple("系统码" if column == "商品编码" else column for column in columns)
            )
        }
        if not matched_kinds:
            raise HTTPException(status_code=400, detail="模板表头不匹配，请使用通用导入模板，不要修改列名或顺序")
        documents: OrderedDict[str, TemplateDocument] = OrderedDict()
        kind = None
        for row_number, values in enumerate(rows, start=2):
            if not any(_text(cell) for cell in values):
                continue
            data = dict(zip(headers, values))
            document_type = _text(data["单据类型"])
            row_kind = next((key for key in matched_kinds if document_type in TEMPLATE_TYPES[key]), None)
            if row_kind is None or kind is not None and row_kind != kind:
                raise HTTPException(status_code=400, detail=f"Excel 第 {row_number} 行：单据类型与模板不匹配：{document_type}")
            kind = row_kind
            document_date = _date(data.get("单据日期", data.get("日期")), datemode)
            if row_kind == "transfer":
                supplier = _text(data.get("出货仓库"))
                warehouse = _text(data.get("入货仓库"))
            else:
                supplier = _text(data.get("单位全名", data.get("往来单位全名")))
                warehouse = _text(data.get("仓库全名"))
            handler = _text(data.get("制单人", data.get("经手人")))
            summary = _text(data.get("摘要"))
            required_fields = {"摘要": summary, "经手人/制单人": handler}
            if kind == "transfer":
                required_fields.update({"出货仓库": supplier, "入货仓库": warehouse})
            elif kind not in {"stock_loss", "stock_gain"}:
                required_fields["往来单位"] = supplier
            if kind not in {"accounting", "transfer"}:
                required_fields["仓库全名"] = warehouse
            for field, value in required_fields.items():
                if not value:
                    raise HTTPException(status_code=400, detail=f"Excel 第 {row_number} 行：{field}不能为空")
            if kind == "accounting":
                subject = _text(data["费用项目名"])
                if not subject:
                    raise HTTPException(status_code=400, detail=f"Excel 第 {row_number} 行：费用项目名不能为空")
                detail = {"product_name": subject, "amount": _amount(data["总金额"], row_number, "总金额")}
            else:
                code = _text(data.get("商品编码", data.get("系统码")))
                if not code:
                    raise HTTPException(status_code=400, detail=f"Excel 第 {row_number} 行：商品编码不能为空")
                detail = {
                    "product_code": code,
                    "quantity": _amount(data.get("销售数量", data.get("数量")), row_number, "数量"),
                    "unit_price": _amount(data.get("单价"), row_number, "单价", required=False),
                }
            key = "\x1f".join((document_date, document_type, supplier, warehouse, handler, summary))
            document = documents.get(key)
            if document is None:
                document = TemplateDocument(key, document_date, document_type, supplier, warehouse, handler, summary, [])
                documents[key] = document
            document.rows.append(detail)
        if not documents:
            raise HTTPException(status_code=400, detail="模板没有可导入的单据")
        assert kind is not None
        for document in documents.values():
            label = f"{document.date} {document.document_type}（摘要：{document.summary}）"
            try:
                validate_template_document_numbers(document.rows, label)
            except ValueError as error:
                raise HTTPException(status_code=400, detail=str(error)) from error
        return kind, sheet_name, list(documents.values())
    except HTTPException:
        raise
    except (ValueError, IndexError, OSError, xlrd.XLRDError) as error:
        raise HTTPException(status_code=400, detail="Excel 文件无法读取") from error
    finally:
        if workbook is not None and hasattr(workbook, "close"):
            workbook.close()
