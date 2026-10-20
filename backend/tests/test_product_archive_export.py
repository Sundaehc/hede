from __future__ import annotations

import io
from datetime import datetime
from decimal import Decimal
from types import SimpleNamespace
from unittest.mock import Mock
from urllib.parse import unquote

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import Integer, MetaData, create_engine, select
from sqlalchemy.pool import StaticPool

from api.routes import import_export, suppliers
from domain import excluded_skus
from domain.color_barcode_schema import COLOR_BARCODE_TABLE
from domain.gj_schema import GJ_MERGED_PRODUCT_INFO_TABLE
from domain.inventory_schema import SUPPLIER_BRAND_TABLE
from domain.schema import PRODUCT_ARCHIVE_TABLES, build_product_archive_table
from domain.size_group_schema import SIZE_GROUP_ITEMS_TABLE, SIZE_GROUPS_TABLE
from domain.vip_schema import JST_PRODUCT_PROFILE_TABLE
from storage.product_repository import ProductRepository
from storage import inventory_repository as inventory_storage


EXPORT_DATE = "2026-09-20"


@pytest.fixture(params=[
    ("supplier_brand_33aa8a0c3b6a", "千百度赫锦"),
    ("supplier_brand_9077c69bd02f", "新增测试品牌乙"),
    ("future_archive_brand", "新增测试品牌丙"),
], ids=["existing-custom", "new-custom", "arbitrary-code"])
def archive_export_client(monkeypatch, request):
    brand, label = request.param
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    metadata = MetaData()
    tables = {
        code: build_product_archive_table(table.name, metadata=metadata)
        for code, table in PRODUCT_ARCHIVE_TABLES.items()
    }
    tables[brand] = build_product_archive_table("manual_product_archive_4960", metadata=metadata)
    for source in (
        SUPPLIER_BRAND_TABLE,
        SIZE_GROUPS_TABLE,
        SIZE_GROUP_ITEMS_TABLE,
        COLOR_BARCODE_TABLE,
        JST_PRODUCT_PROFILE_TABLE,
        GJ_MERGED_PRODUCT_INFO_TABLE,
    ):
        source.to_metadata(metadata)
    brand_table = metadata.tables[SUPPLIER_BRAND_TABLE.name]
    brand_table.c.id.type = Integer()
    brand_table.c.created_at.server_default = None
    brand_table.c.updated_at.server_default = None
    brand_table.c.updated_at.onupdate = None
    monkeypatch.setattr(inventory_storage, "SUPPLIER_BRAND_TABLE", brand_table)
    for table in metadata.tables.values():
        for column in table.columns:
            if column.name in {"created_at", "updated_at"}:
                column.server_default = None
    metadata.create_all(engine)
    with engine.begin() as connection:
        connection.execute(SUPPLIER_BRAND_TABLE.insert(), [
            {"id": 4960, "code": brand, "name": label, "product_archive_enabled": True,
             "product_table_name": "manual_product_archive_4960"},
            {"id": 1, "code": "eblan", "name": "伊伴男鞋", "product_archive_enabled": True,
             "product_table_name": None},
        ])
        base_item = {
            "source_workbook": "手工录入", "source_sheet": "商品档案", "source_row_number": "1",
            "raw_payload": {}, "extra_fields": {}, "created_at": datetime(2026, 9, 20, 10),
            "updated_at": datetime(2026, 9, 20, 10), "last_imported_at": None, "deleted_at": None,
            "product_name": "女休闲鞋", "supplier_name": "靓典（千百度赫锦）", "color": "黑色",
            "color_code": "01", "size_range": "220-225", "barcode_build_rule": "货号+颜色代码+尺码",
            "upper_material": "本品牌材质", "year": "2026", "factory_sku": "00001", "cost": None,
        }
        connection.execute(tables[brand].insert(), [
            {**base_item, "id": 1, "sku": "HJ-001", "original_sku": "STYLE-001", "cost": Decimal("94.50")},
            {**base_item, "id": 2, "sku": "HJ-002", "original_sku": "STYLE-002", "created_at": datetime(2026, 9, 19),
             "last_imported_at": datetime(2026, 9, 20, 10)},
            {**base_item, "id": 3, "sku": "HJ-DELETED", "original_sku": "STYLE-001",
             "deleted_at": datetime(2026, 9, 21), "updated_at": datetime(2026, 9, 21), "cost": Decimal("999")},
            {**base_item, "id": 4, "sku": "HJ-OLD", "original_sku": "STYLE-OLD", "year": "2025",
             "size_range": None, "created_at": datetime(2026, 9, 18)},
        ])
        connection.execute(tables["eblan"].insert(), {
            **base_item, "id": 1, "sku": "EB-001", "original_sku": "EB-001", "cost": Decimal("80"),
        })
    repository = object.__new__(ProductRepository)
    repository.engine = engine
    repository._manual_product_tables = {}
    repository._color_code_cache = {}
    inventory_repository = object.__new__(inventory_storage.InventoryRepository)
    inventory_repository.engine = engine
    log = Mock()
    costs = Mock(side_effect=lambda engine, items: [item.update(cost=Decimal("777")) for item in items])
    monkeypatch.setattr(import_export, "write_operation_log", log)
    monkeypatch.setattr(suppliers, "write_operation_log", Mock())
    monkeypatch.setattr(import_export, "apply_jst_product_costs", costs)
    monkeypatch.setattr(excluded_skus, "EXCLUDED_SKUS", frozenset())

    def create_archive_table(created):
        archive_table = build_product_archive_table(str(created["product_table_name"]), metadata=MetaData())
        for column in (archive_table.c.created_at, archive_table.c.updated_at):
            column.server_default = None
        archive_table.create(engine)

    ensure_archive = Mock(side_effect=create_archive_table)
    monkeypatch.setattr(repository, "ensure_manual_product_archive", ensure_archive)
    app = FastAPI()
    app.state.repository = repository
    app.state.inventory_repository = inventory_repository
    user = {"department_code": "商品部", "role_code": "product_user"}

    @app.middleware("http")
    async def set_current_user(request, call_next):
        request.state.current_user = user
        return await call_next(request)

    app.include_router(import_export.router)
    app.include_router(suppliers.router)
    try:
        with TestClient(app) as client:
            yield SimpleNamespace(
                client=client, repository=repository, tables=tables, log=log, costs=costs,
                brand=brand, label=label, inventory_repository=inventory_repository,
                ensure_archive=ensure_archive, user=user,
            )
    finally:
        engine.dispose()


def _export(fixture, **params):
    response = fixture.client.get("/export", params={"brand": fixture.brand, **params})
    assert response.status_code == 200, response.text
    workbook = load_workbook(io.BytesIO(response.content), data_only=True)
    try:
        values = list(workbook.active.values)
        return SimpleNamespace(
            rows=[dict(zip(values[0], row)) for row in values[1:]],
            sheets=workbook.sheetnames,
            filename=unquote(response.headers["content-disposition"]),
        )
    finally:
        workbook.close()


@pytest.mark.parametrize("params, expected", [
    ({}, ["HJ-OLD", "HJ-002", "HJ-001"]),
    ({"ids": "1,3"}, ["HJ-001"]),
    ({"query": "HJ-002"}, ["HJ-002"]),
    ({"sku_prefix": "HJ-00", "year": "2026"}, ["HJ-002", "HJ-001"]),
    ({"activity_date": EXPORT_DATE}, ["HJ-002", "HJ-001"]),
    ({"activity_date_start": EXPORT_DATE, "activity_date_end": EXPORT_DATE}, ["HJ-002", "HJ-001"]),
    ({"activity_date": EXPORT_DATE, "ids": "4", "query": "not-found", "year": "2025"}, ["HJ-002", "HJ-001"]),
])
def test_custom_brand_standard_exports_use_managed_name_and_active_scope(archive_export_client, params, expected):
    result = _export(archive_export_client, **params)
    sku_header = import_export.EXPORT_LABELS["sku"]
    assert [row[sku_header] for row in result.rows] == expected
    assert result.sheets == [archive_export_client.label]
    assert archive_export_client.label in result.filename
    assert archive_export_client.brand not in result.filename
    assert archive_export_client.log.call_args.kwargs["after_data"]["brand_label"] == archive_export_client.label
    archive_export_client.costs.assert_not_called()


@pytest.mark.parametrize("params, expected", [
    ({}, ["HJ-OLD", "HJ-002", "HJ-001"]),
    ({"ids": "1,3"}, ["HJ-001"]),
    ({"query": "HJ-002"}, ["HJ-002"]),
    ({"sku_prefix": "HJ-00", "year": "2026"}, ["HJ-002", "HJ-001"]),
    ({"activity_date": EXPORT_DATE}, ["HJ-002", "HJ-001"]),
])
def test_custom_brand_price_exports_use_managed_name_and_active_scope(archive_export_client, params, expected):
    result = _export(archive_export_client, mode="price", **params)
    assert [row["货号"] for row in result.rows] == expected
    assert result.sheets == ["物价"]
    assert f"{archive_export_client.label}物价.xlsx" in result.filename
    for row in result.rows:
        assert row["预设售价"] == (94.5 if row["货号"] == "HJ-001" else None)
    archive_export_client.costs.assert_not_called()


@pytest.mark.parametrize("params, expected_styles, count", [
    ({}, {"STYLE-001", "STYLE-002", "STYLE-OLD"}, 5),
    ({"ids": "1,3"}, {"STYLE-001"}, 2),
    ({"query": "HJ-002"}, {"STYLE-002"}, 2),
    ({"sku_prefix": "HJ-00", "year": "2026"}, {"STYLE-001", "STYLE-002"}, 4),
    ({"activity_date": EXPORT_DATE}, {"STYLE-001", "STYLE-002"}, 4),
    ({"activity_date_start": EXPORT_DATE, "activity_date_end": EXPORT_DATE}, {"STYLE-001", "STYLE-002"}, 4),
])
def test_custom_brand_size_exports_preserve_brand_sizes_and_archive_values(archive_export_client, params, expected_styles, count):
    result = _export(archive_export_client, mode="with_sizes", **params)
    assert len(result.rows) == count
    assert {row["款式编码"] for row in result.rows} == expected_styles
    assert {row["品牌"] for row in result.rows} == {archive_export_client.label}
    assert {row["鞋面材质"] for row in result.rows} == {"本品牌材质"}
    assert result.sheets == [f"{archive_export_client.label}带尺码"]
    assert archive_export_client.label in result.filename
    for row in result.rows:
        if row["款式编码"] == "STYLE-001":
            assert row["成本价"] == "94.50"
            assert row["商品编码"] in {"HJ-00101220", "HJ-00101225"}
        else:
            assert row["成本价"] is None
    archive_export_client.costs.assert_not_called()


@pytest.mark.parametrize("params", [{}, {"activity_date": EXPORT_DATE}])
def test_overview_exports_include_new_brand_name_and_exclude_deleted_products(archive_export_client, params):
    result = _export(archive_export_client, brand="all", **params)
    assert {row["品牌"] for row in result.rows} == {archive_export_client.label, "伊伴"}
    assert "HJ-DELETED" not in {row[import_export.EXPORT_LABELS["sku"]] for row in result.rows}
    assert len(result.rows) == (3 if params else 4)
    archive_export_client.costs.assert_not_called()


@pytest.mark.parametrize("params, expected", [
    ({"query": "HJ-002"}, {"HJ-002"}),
    ({"sku_prefix": "HJ-00", "year": "2026"}, {"HJ-001", "HJ-002"}),
])
def test_overview_search_includes_custom_brands(archive_export_client, params, expected):
    result = _export(archive_export_client, brand="all", **params)
    assert {row[import_export.EXPORT_LABELS["sku"]] for row in result.rows} == expected
    assert {row["品牌"] for row in result.rows} == {archive_export_client.label}


def test_all_brand_price_export_includes_new_brand(archive_export_client):
    result = _export(archive_export_client, brand="all", mode="price", year="2026")
    assert {row["货号"] for row in result.rows} == {"HJ-001", "HJ-002", "EB-001"}


def test_custom_brand_size_export_does_not_use_other_brand_external_data(archive_export_client, monkeypatch):
    foreign_profiles = Mock(return_value=[{
        "id": 999, "product_code": "FOREIGN-SIZE", "style_code": "STYLE-OLD", "size_barcode": "999",
        "raw_payload": {"品牌": "其他品牌", "成本价": "999"},
    }])
    foreign_goods = Mock(return_value={"STYLE-OLD": {"brand": "其他品牌", "upper_material": "其他材质"}})
    monkeypatch.setattr(import_export, "_load_product_profile_rows", foreign_profiles)
    monkeypatch.setattr(import_export, "_load_gj_rows", foreign_goods)
    result = _export(archive_export_client, mode="with_sizes", ids="4")
    assert len(result.rows) == 1
    assert result.rows[0]["商品编码"] == "HJ-OLD"
    assert result.rows[0]["尺码条码"] is None
    assert result.rows[0]["品牌"] == archive_export_client.label
    assert result.rows[0]["鞋面材质"] == "本品牌材质"
    foreign_profiles.assert_not_called()
    foreign_goods.assert_not_called()


def test_custom_brand_size_export_uses_maintained_size_group(archive_export_client):
    with archive_export_client.repository.engine.begin() as connection:
        connection.execute(SIZE_GROUPS_TABLE.insert().values(id=1, name="女鞋尺码组"))
        connection.execute(SIZE_GROUP_ITEMS_TABLE.insert(), [
            {"id": 1, "size_group_id": 1, "size_name": "34", "barcode": "220", "sort_order": 1},
            {"id": 2, "size_group_id": 1, "size_name": "35", "barcode": "225", "sort_order": 2},
        ])
        connection.execute(
            archive_export_client.tables[archive_export_client.brand].update()
            .where(archive_export_client.tables[archive_export_client.brand].c.id == 1)
            .values(size_range="女鞋尺码组", updated_at=datetime(2026, 9, 20))
        )
    result = _export(archive_export_client, mode="with_sizes", ids="1")
    assert {row["商品编码"] for row in result.rows} == {"HJ-00101220", "HJ-00101225"}
    assert {row["品牌"] for row in result.rows} == {archive_export_client.label}


def test_size_export_of_only_deleted_selection_reports_no_products(archive_export_client):
    response = archive_export_client.client.get("/export", params={
        "brand": archive_export_client.brand, "mode": "with_sizes", "ids": "3",
    })
    assert response.status_code == 404
    archive_export_client.log.assert_not_called()


@pytest.mark.parametrize("params", [{"query": "not-found"}, {"sku_prefix": "not-found"}])
@pytest.mark.parametrize("mode", [None, "price", "with_sizes"])
def test_custom_brand_exports_with_no_matches_produce_empty_workbook(archive_export_client, params, mode):
    result = _export(archive_export_client, **params, **({"mode": mode} if mode else {}))
    assert result.rows == []


@pytest.mark.parametrize("mode, suffix", [(None, ""), ("with_sizes", "带尺码"), ("price", "物价")])
def test_exports_follow_brand_renames_without_restart(archive_export_client, mode, suffix):
    renamed_label = "赫锦新名称"
    with archive_export_client.repository.engine.begin() as connection:
        connection.execute(
            SUPPLIER_BRAND_TABLE.update()
            .where(SUPPLIER_BRAND_TABLE.c.code == archive_export_client.brand)
            .values(name=renamed_label, updated_at=datetime(2026, 9, 20))
        )
    params = {"mode": mode} if mode else {}
    result = _export(archive_export_client, **params)
    assert renamed_label in result.filename
    assert result.sheets == (["物价"] if mode == "price" else [f"{renamed_label}{suffix}"])


@pytest.mark.parametrize("mode", [None, "price", "with_sizes"])
def test_new_brand_export_preflight_accepts_each_mode(archive_export_client, mode):
    params = {"brand": archive_export_client.brand}
    if mode:
        params["mode"] = mode
    response = archive_export_client.client.head("/export", params=params)
    assert response.status_code == 200
    assert response.content == b""
    archive_export_client.log.assert_not_called()


@pytest.mark.parametrize("mode", [None, "with_sizes"])
def test_custom_brand_names_are_safe_for_excel_sheet_titles(archive_export_client, mode):
    label = "千百度[赫锦]/新品:*?\\" + "长名称" * 15
    with archive_export_client.repository.engine.begin() as connection:
        connection.execute(
            SUPPLIER_BRAND_TABLE.update()
            .where(SUPPLIER_BRAND_TABLE.c.code == archive_export_client.brand)
            .values(name=label, updated_at=datetime(2026, 9, 20))
        )
    result = _export(archive_export_client, **({"mode": mode} if mode else {}))
    assert len(result.sheets[0]) <= 31
    assert not set(result.sheets[0]) & set("[]:*?/\\")
    assert label in result.filename


def test_builtin_brand_still_backfills_missing_export_costs(archive_export_client):
    with archive_export_client.repository.engine.begin() as connection:
        connection.execute(
            archive_export_client.tables["eblan"].update()
            .values(cost=None, updated_at=datetime(2026, 9, 20))
        )
    result = _export(archive_export_client, brand="eblan", activity_date=EXPORT_DATE)
    assert result.sheets == ["伊伴"]
    assert result.rows[0][import_export.EXPORT_LABELS["cost"]] == 777
    archive_export_client.costs.assert_called_once()


def test_builtin_brand_size_export_retains_external_profile_lookup(archive_export_client, monkeypatch):
    with archive_export_client.repository.engine.begin() as connection:
        connection.execute(
            archive_export_client.tables["eblan"].update()
            .values(size_range=None, updated_at=datetime(2026, 9, 20))
        )
    profile_lookup = Mock(return_value=[{
        "id": 1, "product_code": "EB-00134", "style_code": "EB-001", "color_name": "黑色", "size_barcode": "34",
        "raw_payload": {"品牌": "伊伴"},
    }])
    goods_lookup = Mock(return_value={})
    monkeypatch.setattr(import_export, "_load_product_profile_rows", profile_lookup)
    monkeypatch.setattr(import_export, "_load_gj_rows", goods_lookup)
    result = _export(archive_export_client, brand="eblan", mode="with_sizes")
    assert len(result.rows) == 1
    assert result.rows[0]["商品编码"] == "EB-00134"
    assert result.rows[0]["品牌"] == "伊伴"
    profile_lookup.assert_called_once()
    goods_lookup.assert_called_once()


def test_brands_created_after_initial_exports_are_discovered_without_restart(archive_export_client):
    initial = _export(archive_export_client, brand="all")
    assert {row["品牌"] for row in initial.rows} == {archive_export_client.label, "伊伴"}
    new_brands = []
    for index, label in enumerate(("未来品牌甲", "未来品牌乙"), start=1):
        response = archive_export_client.client.post("/supplier-brands", json={"name": label})
        assert response.status_code == 200, response.text
        created = response.json()["item"]
        archive_export_client.ensure_archive.assert_called_with(created)
        brand = str(created["code"])
        assert brand.startswith("supplier_brand_")
        assert brand not in import_export.BRAND_LABELS
        table = archive_export_client.repository._table_for_brand(brand)
        assert archive_export_client.repository.is_product_archive_brand(brand)
        assert brand in archive_export_client.repository.product_archive_brands()
        for mode in (None, "price", "with_sizes"):
            params = {"mode": mode} if mode else {}
            empty = _export(archive_export_client, brand=brand, **params)
            assert empty.rows == []
            assert label in empty.filename
        with archive_export_client.repository.engine.begin() as connection:
            connection.execute(table.insert(), [
                {
                    "id": 1, "sku": "SHARED-SKU", "original_sku": "SHARED-STYLE",
                    "source_workbook": "手工录入", "source_sheet": "商品档案", "source_row_number": "1",
                    "raw_payload": {"品牌": "旧品牌"}, "created_at": datetime(2026, 9, 20, 10),
                    "updated_at": datetime(2026, 9, 20, 10), "deleted_at": None,
                    "supplier_name": f"{label}供应商", "upper_material": f"{label}材质",
                    "size_range": "220-225", "color": "黑色", "color_code": "01",
                    "cost": Decimal(index * 10), "year": "2026",
                },
                {
                    "id": 2, "sku": "SHARED-DELETED", "original_sku": "SHARED-STYLE",
                    "source_workbook": "手工录入", "source_sheet": "商品档案", "source_row_number": "2",
                    "raw_payload": {}, "created_at": datetime(2026, 9, 20, 10),
                    "updated_at": datetime(2026, 9, 21, 10), "deleted_at": datetime(2026, 9, 21, 10),
                    "supplier_name": "已删除供应商", "upper_material": "已删除材质",
                    "size_range": "220-225", "color": "黑色", "color_code": "01",
                    "cost": Decimal("999"), "year": "2026",
                },
            ])
        new_brands.append((created, index))

    for created, index in new_brands:
        brand, label = str(created["code"]), str(created["name"])
        for mode in (None, "price", "with_sizes"):
            for filters in (
                {}, {"ids": "1,2"}, {"query": "SHARED-SKU"},
                {"sku_prefix": "SHARED", "year": "2026"}, {"activity_date": EXPORT_DATE},
            ):
                params = {"brand": brand, **filters, **({"mode": mode} if mode else {})}
                assert archive_export_client.client.head("/export", params=params).status_code == 200
                result = _export(archive_export_client, **params)
                assert len(result.rows) == (2 if mode == "with_sizes" else 1)
                assert label in result.filename
                assert brand not in result.filename
                if mode == "with_sizes":
                    assert {row["品牌"] for row in result.rows} == {label}
                    assert {row["供应商名"] for row in result.rows} == {f"{label}供应商"}
                    assert {row["鞋面材质"] for row in result.rows} == {f"{label}材质"}
                    assert {row["成本价"] for row in result.rows} == {f"{index * 10}.00"}
                    assert {row["商品编码"] for row in result.rows} == {"SHARED-SKU01220", "SHARED-SKU01225"}
                elif mode == "price":
                    assert result.rows[0]["主供应商"] == f"{label}供应商"
                    assert result.rows[0]["预设售价"] == index * 10
                else:
                    assert result.rows[0][import_export.EXPORT_LABELS["supplier_name"]] == f"{label}供应商"
                    assert result.rows[0][import_export.EXPORT_LABELS["cost"]] == index * 10

    for mode in (None, "price"):
        for filters in ({}, {"activity_date": EXPORT_DATE}):
            result = _export(archive_export_client, brand="all", query="SHARED", **filters, **({"mode": mode} if mode else {}))
            sku_header = "货号" if mode == "price" else import_export.EXPORT_LABELS["sku"]
            shared_rows = [row for row in result.rows if row[sku_header] == "SHARED-SKU"]
            assert len(shared_rows) == 2
            assert "SHARED-DELETED" not in {row[sku_header] for row in result.rows}
            if mode is None:
                assert {row["品牌"] for row in shared_rows} == {item["name"] for item, _ in new_brands}

    created = new_brands[0][0]
    renamed_label = "未来品牌改名"
    response = archive_export_client.client.put(f"/supplier-brands/{created['id']}", json={"name": renamed_label})
    assert response.status_code == 200, response.text
    for mode in (None, "price", "with_sizes"):
        result = _export(archive_export_client, brand=created["code"], **({"mode": mode} if mode else {}))
        assert renamed_label in result.filename
        if mode == "with_sizes":
            assert {row["品牌"] for row in result.rows} == {renamed_label}
    renamed_overview = _export(archive_export_client, brand="all", query="SHARED-SKU")
    assert {row["品牌"] for row in renamed_overview.rows} == {renamed_label, new_brands[1][0]["name"]}
    archive_export_client.costs.assert_not_called()


def test_custom_brand_export_does_not_mutate_product_data(archive_export_client):
    table = archive_export_client.tables[archive_export_client.brand]
    with archive_export_client.repository.engine.connect() as connection:
        before = [dict(row) for row in connection.execute(select(table).order_by(table.c.id)).mappings()]
    for mode in (None, "price", "with_sizes"):
        for filters in ({}, {"activity_date": EXPORT_DATE}):
            _export(archive_export_client, **filters, **({"mode": mode} if mode else {}))
    with archive_export_client.repository.engine.connect() as connection:
        after = [dict(row) for row in connection.execute(select(table).order_by(table.c.id)).mappings()]
    assert after == before


@pytest.mark.parametrize("mode", [None, "price", "with_sizes"])
def test_blank_and_zero_costs_are_not_replaced_for_new_brands(archive_export_client, mode):
    table = archive_export_client.tables[archive_export_client.brand]
    with archive_export_client.repository.engine.begin() as connection:
        connection.execute(table.update().where(table.c.id == 1).values(cost=0, updated_at=datetime(2026, 9, 20)))
    params = {"mode": mode} if mode else {}
    result = _export(archive_export_client, ids="1,2", **params)
    cost_header = "成本价" if mode == "with_sizes" else "预设售价" if mode == "price" else import_export.EXPORT_LABELS["cost"]
    assert {row[cost_header] for row in result.rows} == ({"0.00", None} if mode == "with_sizes" else {0, None})
    archive_export_client.costs.assert_not_called()


@pytest.mark.parametrize("department", ["客服部", "美工部"])
@pytest.mark.parametrize("mode", [None, "with_sizes"])
@pytest.mark.parametrize("dated", [False, True])
def test_new_brand_exports_keep_cost_visibility_restrictions(archive_export_client, department, mode, dated):
    archive_export_client.user["department_code"] = department
    params = {"mode": mode} if mode else {}
    if dated:
        params["activity_date"] = EXPORT_DATE
    result = _export(archive_export_client, **params)
    assert result.rows
    headers = set(result.rows[0])
    assert "成本价" not in headers
    assert import_export.EXPORT_LABELS["cost"] not in headers
    if mode is None:
        overview = _export(archive_export_client, brand="all", **params)
        assert "成本价" not in overview.rows[0]
        assert import_export.EXPORT_LABELS["cost"] not in overview.rows[0]
    archive_export_client.costs.assert_not_called()


@pytest.mark.parametrize("department", ["客服部", "美工部"])
@pytest.mark.parametrize("method", ["get", "head"])
def test_new_brand_price_export_denies_cost_restricted_users(archive_export_client, department, method):
    archive_export_client.user["department_code"] = department
    response = getattr(archive_export_client.client, method)("/export", params={
        "brand": archive_export_client.brand, "mode": "price",
    })
    assert response.status_code == 403
    archive_export_client.log.assert_not_called()
