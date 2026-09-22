from __future__ import annotations

import asyncio
import io
from types import SimpleNamespace
from unittest.mock import AsyncMock, Mock

import pytest
from fastapi import HTTPException, UploadFile
from openpyxl import Workbook
from sqlalchemy import create_engine, event

from api.routes import inventory as inventory_routes
from domain.inventory_schema import SUPPLIER_TABLE
from storage.inventory_repository import InventoryRepository


SUPPLIER_NAME = "友宝保罗（千百度）"
PURCHASE_HEADERS = [
    "供应商", "商品编码", "数量", "采购单备注", "采购日期",
    "协议到货日期", "收货仓库", "经办人",
]
GENERIC_HEADERS = [
    "supplier", "product_code", "quantity", "summary", "date",
    "delivery_date", "warehouse", "handler", "document_type",
]


@pytest.fixture
def supplier_repository():
    engine = create_engine("sqlite://")
    event.listen(
        engine,
        "connect",
        lambda connection, _record: connection.create_function(
            "date_trunc", 2, lambda _unit, value: value,
        ),
    )
    SUPPLIER_TABLE.create(engine)
    with engine.begin() as connection:
        connection.execute(SUPPLIER_TABLE.insert(), [
            {"id": 1, "brand": "cbanner_mens", "name": SUPPLIER_NAME},
            {"id": 2, "brand": "cbanner_mens", "name": "Acme工厂"},
        ])
    repository = Mock(spec=InventoryRepository)
    repository.engine = engine
    repository.get_supplier_by_name.side_effect = lambda name: InventoryRepository.get_supplier_by_name(
        SimpleNamespace(engine=engine), name,
    )
    repository.get_record_for_append.return_value = None
    repository.create_record.side_effect = lambda payload: {"id": 1, **payload}
    repository.get_record.return_value = {
        "id": 1, "document_type": "进货订单", "supplier": SUPPLIER_NAME,
    }
    repository.list_details.return_value = []
    repository.merge_imported_details.return_value = {"added": 1, "updated": 0}
    try:
        yield repository
    finally:
        engine.dispose()


@pytest.fixture
def detail_builder(monkeypatch):
    builder = Mock(return_value=[{
        "product_code": "C5563406D8080", "quantity": "6", "amount": "60",
    }])
    monkeypatch.setattr(inventory_routes, "_build_purchase_details_from_rows", builder)
    monkeypatch.setattr(inventory_routes, "_build_purchase_detail_lookup", lambda *_args: {})
    monkeypatch.setattr(inventory_routes, "write_operation_log", lambda *_args, **_kwargs: None)
    return builder


def _request(repository, document_type="进货订单"):
    return SimpleNamespace(
        app=SimpleNamespace(state=SimpleNamespace(inventory_repository=repository)),
        form=AsyncMock(return_value={"document_type": document_type}),
    )


def _upload(headers, rows, *, title=False):
    workbook = Workbook()
    worksheet = workbook.active
    if title:
        worksheet.append(["采购单明细"])
    worksheet.append(headers)
    for row in rows:
        worksheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    buffer.seek(0)
    return UploadFile(filename="采购单.xlsx", file=buffer)


def _purchase_row(supplier=SUPPLIER_NAME, *, summary="测试采购单"):
    return [supplier, "C5563406D8080240", 6, summary, "2026-09-20", "2026-10-01", "测试仓库", "测试员"]


def _assert_no_writes(repository):
    repository.create_record.assert_not_called()
    repository.create_details.assert_not_called()
    repository.create_supplier.assert_not_called()
    repository.merge_imported_details.assert_not_called()
    repository.update_record.assert_not_called()


@pytest.mark.parametrize("supplier", [SUPPLIER_NAME, "Acme工厂"])
@pytest.mark.parametrize("append", [False, True])
def test_purchase_import_accepts_exact_supplier(supplier_repository, detail_builder, supplier, append):
    if append:
        supplier_repository.get_record_for_append.return_value = {"id": 1, "supplier": supplier}
    result = asyncio.run(inventory_routes.import_purchase_inventory(
        _request(supplier_repository), _upload(PURCHASE_HEADERS, [_purchase_row(supplier)]),
    ))

    assert result["created"] == (0 if append else 1)
    assert result["appended"] == (1 if append else 0)
    supplier_repository.create_details.assert_called_once()
    supplier_repository.create_supplier.assert_not_called()


@pytest.mark.parametrize("supplier", [
    "不存在的工厂", "友宝保罗", "友宝保罗(千百度)", "acme工厂", "友宝保罗 （千百度）",
])
def test_purchase_import_rejects_non_exact_supplier(supplier_repository, detail_builder, supplier):
    with pytest.raises(HTTPException) as error:
        asyncio.run(inventory_routes.import_purchase_inventory(
            _request(supplier_repository), _upload(PURCHASE_HEADERS, [_purchase_row(supplier)]),
        ))

    assert error.value.status_code == 400
    assert "Excel 第 2 行" in error.value.detail
    assert supplier in error.value.detail
    assert "完全一致" in error.value.detail
    detail_builder.assert_not_called()
    _assert_no_writes(supplier_repository)


@pytest.mark.parametrize("summary", ["测试采购单", "另一张采购单"])
@pytest.mark.parametrize("append", [False, True])
def test_purchase_import_validates_later_rows_before_any_writes(
    supplier_repository, detail_builder, summary, append,
):
    if append:
        supplier_repository.get_record_for_append.return_value = {"id": 1}
    rows = [_purchase_row(), [], _purchase_row("不存在的工厂", summary=summary)]
    with pytest.raises(HTTPException) as error:
        asyncio.run(inventory_routes.import_purchase_inventory(
            _request(supplier_repository), _upload(PURCHASE_HEADERS, rows, title=True),
        ))

    assert "Excel 第 5 行" in error.value.detail
    assert "本次未导入任何数据" in error.value.detail
    detail_builder.assert_not_called()
    _assert_no_writes(supplier_repository)


def test_purchase_import_carries_supplier_through_merged_cells(supplier_repository, detail_builder):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(PURCHASE_HEADERS)
    worksheet.append(_purchase_row())
    worksheet.append(_purchase_row())
    worksheet.merge_cells("A2:A3")
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    buffer.seek(0)

    result = asyncio.run(inventory_routes.import_purchase_inventory(
        _request(supplier_repository), UploadFile(filename="采购单.xlsx", file=buffer),
    ))

    assert result["created"] == 1
    assert supplier_repository.create_record.call_args.args[0]["supplier"] == SUPPLIER_NAME
    assert len(detail_builder.call_args.args[1]) == 2


@pytest.mark.parametrize("rows", [
    [_purchase_row("")],
    [_purchase_row(""), _purchase_row()],
])
def test_purchase_import_rejects_missing_supplier(supplier_repository, detail_builder, rows):
    with pytest.raises(HTTPException) as error:
        asyncio.run(inventory_routes.import_purchase_inventory(
            _request(supplier_repository), _upload(PURCHASE_HEADERS, rows),
        ))

    assert error.value.status_code == 400
    assert "供应商" in error.value.detail
    _assert_no_writes(supplier_repository)


@pytest.mark.parametrize("supplier", ["不存在的工厂", "友宝保罗(千百度)", "acme工厂", ""])
def test_generic_import_cannot_bypass_purchase_supplier_validation(
    supplier_repository, detail_builder, supplier,
):
    rows = [
        [*_purchase_row(), "进货订单"],
        [*_purchase_row(supplier), "进货订单"],
    ]
    with pytest.raises(HTTPException) as error:
        asyncio.run(inventory_routes.import_inventory(
            _request(supplier_repository), _upload(GENERIC_HEADERS, rows),
        ))

    assert error.value.status_code == 400
    assert "Excel 第 3 行" in error.value.detail
    _assert_no_writes(supplier_repository)


def test_generic_import_accepts_exact_purchase_supplier(supplier_repository, detail_builder):
    result = asyncio.run(inventory_routes.import_inventory(
        _request(supplier_repository),
        _upload(GENERIC_HEADERS, [[*_purchase_row(), "进货订单"]]),
    ))

    assert result["created"] == 1
    assert result["skipped"] == 0
    supplier_repository.create_details.assert_called_once()
    supplier_repository.create_supplier.assert_not_called()


def test_generic_non_purchase_import_keeps_supplier_creation(supplier_repository, detail_builder):
    result = asyncio.run(inventory_routes.import_inventory(
        _request(supplier_repository),
        _upload(GENERIC_HEADERS, [[*_purchase_row("新供应商"), "进货单"]]),
    ))

    assert result["created"] == 1
    supplier_repository.create_supplier.assert_called_once_with({"name": "新供应商"})


def test_non_purchase_detail_import_keeps_optional_supplier(supplier_repository, detail_builder):
    supplier_repository.get_warehouse_by_name.return_value = {"brand": "cbanner_mens"}
    result = asyncio.run(inventory_routes.import_purchase_inventory(
        _request(supplier_repository, "报溢单"),
        _upload(PURCHASE_HEADERS, [_purchase_row("")]),
    ))

    assert result["created"] == 1
    supplier_repository.get_supplier_by_name.assert_not_called()


@pytest.mark.parametrize("record_supplier, row_supplier", [
    ("不存在的工厂", None),
    ("", None),
    (SUPPLIER_NAME, "友宝保罗(千百度)"),
])
def test_purchase_reimport_validates_record_and_excel_suppliers(
    supplier_repository, detail_builder, record_supplier, row_supplier,
):
    supplier_repository.get_record.return_value["supplier"] = record_supplier
    with pytest.raises(HTTPException) as error:
        asyncio.run(inventory_routes.reimport_inventory_details_from_excel(
            _request(supplier_repository), 1,
            _upload(PURCHASE_HEADERS, [_purchase_row(row_supplier)]),
        ))

    assert error.value.status_code == 400
    detail_builder.assert_not_called()
    _assert_no_writes(supplier_repository)


def test_purchase_reimport_uses_record_supplier_for_detail_only_file(supplier_repository, detail_builder):
    result = asyncio.run(inventory_routes.reimport_inventory_details_from_excel(
        _request(supplier_repository), 1,
        _upload(["商品编码", "数量"], [["C5563406D8080240", 6]]),
    ))

    assert result["added"] == 1
    supplier_repository.merge_imported_details.assert_called_once()


def test_legacy_xls_parser_preserves_excel_row_numbers(monkeypatch):
    values = [["采购单明细"], PURCHASE_HEADERS, [], _purchase_row()]
    worksheet = SimpleNamespace(nrows=len(values), name="采购单", row_values=values.__getitem__)
    workbook = SimpleNamespace(nsheets=1, sheet_by_index=lambda _index: worksheet)
    monkeypatch.setattr(inventory_routes.xlrd, "open_workbook", lambda **_kwargs: workbook)

    rows, sheet_name = inventory_routes._read_purchase_import_rows_xls(b"xls")

    assert sheet_name == "采购单"
    assert rows[0]["source_row_number"] == 4


def test_purchase_template_explains_exact_supplier_matching():
    workbook = inventory_routes._build_purchase_order_import_template()
    try:
        assert "供应商管理中的完整名称完全一致" in workbook.active.cell(2, 1).value
        assert "不会自动新增供应商" in workbook.active.cell(2, 1).value
    finally:
        workbook.close()
