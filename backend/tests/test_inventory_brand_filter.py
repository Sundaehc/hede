from __future__ import annotations

import asyncio
from datetime import date, datetime
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import Mock

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine, event

from api.routes.inventory import export_inventory, export_purchase_inbound_details, list_inventory_brands, list_purchase_inbound_details
from domain.inventory_brands import inventory_brand_filter_aliases, inventory_brand_options
from domain.inventory_schema import (
    GENERAL_CUSTOMER_BRAND_TABLE, GENERAL_CUSTOMER_SHOP_TABLE, GENERAL_CUSTOMER_UNIT_TABLE,
    INVENTORY_DETAIL_TABLE, INVENTORY_TABLE, SUPPLIER_BRAND_TABLE, SUPPLIER_TABLE, WAREHOUSE_BRAND_TABLE, WAREHOUSE_TABLE,
)
from storage.inventory_repository import InventoryRepository


@pytest.fixture
def brand_repository():
    engine = create_engine("sqlite://")
    event.listen(engine, "connect", lambda connection, record: connection.create_function(
        "date_trunc", 2, lambda unit, value: value,
    ))
    for table in (
        INVENTORY_TABLE, INVENTORY_DETAIL_TABLE, SUPPLIER_BRAND_TABLE, SUPPLIER_TABLE, WAREHOUSE_BRAND_TABLE,
        WAREHOUSE_TABLE, GENERAL_CUSTOMER_BRAND_TABLE, GENERAL_CUSTOMER_SHOP_TABLE, GENERAL_CUSTOMER_UNIT_TABLE,
    ):
        table.create(engine)
    with engine.begin() as connection:
        connection.execute(SUPPLIER_BRAND_TABLE.insert(), [
            {"id": 1, "code": "yandou", "name": "烟斗"},
            {"id": 2, "code": "cbanner_mens", "name": "千百度男鞋"},
            {"id": 3, "code": "cbanner_womens", "name": "千百度女鞋"},
            {"id": 4, "code": "eblan", "name": "伊伴男鞋"},
            {"id": 5, "code": "eblan_womens", "name": "伊伴女鞋"},
            {"id": 6, "code": "custom", "name": "自定义品牌"},
        ])
        connection.execute(SUPPLIER_TABLE.insert(), [
            {"id": 1, "brand": "yandou", "name": "烟斗供应商"},
            {"id": 2, "brand": "cbanner_mens", "name": "男鞋供应商"},
            {"id": 3, "brand": "cbanner_womens", "name": "女鞋供应商"},
            {"id": 4, "brand": "eblan", "name": "伊伴男供应商"},
            {"id": 5, "brand": "eblan_womens", "name": "伊伴女供应商"},
            {"id": 6, "brand": "custom", "name": "自定义供应商"},
        ])
        connection.execute(WAREHOUSE_BRAND_TABLE.insert(), [
            {"id": 1, "name": "烟斗仓库"}, {"id": 2, "name": "NI仓库"},
        ])
        connection.execute(WAREHOUSE_TABLE.insert(), [
            {"id": 1, "brand": "烟斗仓库", "name": "烟斗总仓"},
            {"id": 2, "brand": "NI仓库", "name": "NI总仓"},
        ])
        connection.execute(GENERAL_CUSTOMER_BRAND_TABLE.insert(), [
            {"id": 1, "name": "烟斗"}, {"id": 2, "name": "千百度"}, {"id": 3, "name": "自定义品牌"},
        ])
        connection.execute(GENERAL_CUSTOMER_SHOP_TABLE.insert(), [
            {"id": 1, "customer_name": "烟斗", "shop_name": "烟斗店铺", "is_active": False},
            {"id": 2, "customer_name": "千百度", "shop_name": "千百度店铺", "is_active": True},
            {"id": 3, "customer_name": "自定义品牌", "shop_name": "自定义店铺", "is_active": True},
            {"id": 4, "customer_name": "烟斗", "shop_name": "烟斗其他店铺", "is_active": True},
        ])
        connection.execute(GENERAL_CUSTOMER_UNIT_TABLE.insert(), [
            {"id": 1, "shop_id": 1, "unit_name": "烟斗单位"},
            {"id": 2, "shop_id": 4, "unit_name": "烟斗单位"},
            {"id": 3, "shop_id": 3, "unit_name": "自定义单位"},
        ])
        records = [
            (1, "进货单", "未知", "烟斗总仓"),
            (2, "应付款增加", "烟斗供应商", None),
            (3, "应收款减少", "烟斗店铺", None),
            (4, "应收款增加", "烟斗单位", None),
            (5, "批发销售单", "烟斗店铺", "NI总仓"),
            (6, "进货单", "未知", "NI总仓"),
            (7, "进货单", "烟斗供应商", "烟斗总仓"),
            (8, "同价调拨单", "烟斗总仓", "NI总仓"),
            (9, "进货单", "烟斗总仓", "NI总仓"),
            (10, "应付款增加", "烟斗供应商", None),
            (11, "应付款增加", "女鞋供应商", None),
            (12, "应收款增加", "千百度店铺", None),
            (13, "应付款增加", "伊伴女供应商", None),
            (14, "应付款增加", "伊伴男供应商", None),
            (15, "应收款增加", "自定义单位", None),
            (16, "应付款增加", "自定义供应商", None),
            (17, "应付款减少", " 烟斗供应商 ", None),
        ]
        connection.execute(INVENTORY_TABLE.insert(), [
            {
                "id": record_id, "document_number": f"DOC-{record_id}", "document_type": document_type,
                "supplier": supplier, "warehouse": warehouse, "date": "2026-10-01",
                "date_value": date(2026, 10, 1), "summary": "烟斗", "total_count": 1, "amount": 10,
                "deleted_at": datetime.now() if record_id == 10 else None,
            }
            for record_id, document_type, supplier, warehouse in records
        ])
    repository = object.__new__(InventoryRepository)
    repository.engine = engine
    repository.purge_expired_deleted_records = lambda: 0
    try:
        yield repository
    finally:
        engine.dispose()


@pytest.mark.parametrize("brand", ["yandou", "烟斗", "烟斗仓库", " YANDOU "])
def test_brand_matches_warehouse_supplier_shop_and_disabled_shop_units(brand_repository, brand):
    result = brand_repository.list_records(brands=[brand], page=1, page_size=100)
    assert {item["id"] for item in result["items"]} == {1, 2, 3, 4, 5, 7, 8, 17}
    assert result["total"] == 8
    assert result["totals"]["all"]["amount"] == "80"


def test_overlapping_brand_matches_do_not_duplicate_records_or_totals(brand_repository):
    result = brand_repository.list_records(brands=["yandou", "烟斗", "烟斗仓库"], page=1, page_size=3)
    assert result["total"] == 8
    assert len(result["items"]) == 3
    assert result["totals"]["all"]["amount"] == "80"
    assert result["totals"]["current_page"]["amount"] == "30"


def test_select_all_brand_scope_matches_list_scope_without_pagination(brand_repository):
    assert set(brand_repository.list_record_ids(brands=["yandou", "烟斗仓库"])) == {1, 2, 3, 4, 5, 7, 8, 17}


@pytest.mark.parametrize("brand,expected_ids", [
    ("千百度仓库", {12}), ("cbanner_mens", {12}), ("千百度男鞋", {12}),
    ("千百度女鞋仓库", {11}), ("伊伴男鞋仓库", {14}), ("eblan_womens", {13}),
    ("自定义品牌", {15, 16}), ("custom", {15, 16}), ("自定义品牌仓库", {15, 16}),
])
def test_gender_and_custom_brand_names_remain_distinct(brand_repository, brand, expected_ids):
    result = brand_repository.list_records(brands=[brand], page=1, page_size=100)
    assert {item["id"] for item in result["items"]} == expected_ids


def test_brand_filter_still_intersects_explicit_warehouse_and_document_filters(brand_repository):
    result = brand_repository.list_records(
        brands=["烟斗"], warehouses=["NI总仓"], document_types=["批发销售单"], page=1, page_size=100,
    )
    assert [item["id"] for item in result["items"]] == [5]
    multiple = brand_repository.list_records(brands=["烟斗", "custom"], page=1, page_size=100)
    assert {item["id"] for item in multiple["items"]} == {1, 2, 3, 4, 5, 7, 8, 15, 16, 17}


def test_brand_options_merge_all_master_data_without_duplicate_aliases(brand_repository):
    options = brand_repository.list_inventory_brand_options()
    assert {item["value"] for item in options} == {
        "yandou", "ni", "cbanner_mens", "cbanner_womens", "eblan", "eblan_womens", "custom",
    }
    assert options.count({"value": "yandou", "label": "烟斗"}) == 1
    assert {"value": "custom", "label": "自定义品牌"} in options
    request = SimpleNamespace(app=SimpleNamespace(state=SimpleNamespace(inventory_repository=brand_repository)))
    assert list_inventory_brands(request) == {"items": options}


def test_brand_options_include_unconfigured_custom_supplier_and_customer_brands():
    options = inventory_brand_options([], [{"code": "new_code", "name": "新品牌"}], ["新品牌", "其他品牌"])
    assert options == [{"value": "new_code", "label": "新品牌"}, {"value": "其他品牌", "label": "其他品牌"}]
    assert "新品牌仓库" in inventory_brand_filter_aliases(["new_code"], [{"code": "new_code", "name": "新品牌"}])


def test_export_uses_the_same_brand_scope_including_warehouse_free_accounting_records(brand_repository):
    brand_repository.list_details_for_documents = Mock(return_value=[])
    request = SimpleNamespace(app=SimpleNamespace(state=SimpleNamespace(inventory_repository=brand_repository)))
    response = export_inventory(request, brands=["烟斗仓库"])

    async def read_content():
        return b"".join([chunk async for chunk in response.body_iterator])

    workbook = load_workbook(BytesIO(asyncio.run(read_content())), data_only=True)
    try:
        rows = list(workbook["经营历程"].iter_rows(values_only=True))
        headers = rows[0]
        numbers = {row[headers.index("单据编号")] for row in rows[1:]}
        assert numbers == {f"DOC-{record_id}" for record_id in (1, 2, 3, 4, 5, 7, 8, 17)}
    finally:
        workbook.close()


def _add_inbound_brand_details(repository):
    with repository.engine.begin() as connection:
        connection.execute(INVENTORY_TABLE.insert(), [
            {"id": 18, "document_number": "DOC-18", "document_type": "进货退货单", "supplier": "未知", "warehouse": "烟斗总仓", "date": "2026-10-02", "date_value": date(2026, 10, 2)},
            {"id": 19, "document_number": "DOC-19", "document_type": "进货单", "supplier": "自定义供应商", "warehouse": "未知", "date": "2026-10-02", "date_value": date(2026, 10, 2)},
            {"id": 20, "document_number": "DOC-20", "document_type": "进货单", "supplier": "女鞋供应商", "warehouse": "未知", "date": "2026-10-02", "date_value": date(2026, 10, 2)},
        ])
        connection.execute(INVENTORY_DETAIL_TABLE.insert(), [
            {"id": record_id, "document_id": record_id, "product_code": "TEST", "product_name": "休闲鞋", "color_name": "黑色", "quantity": 2 if record_id == 18 else 1, "amount": 20 if record_id == 18 else 10}
            for record_id in range(1, 21)
        ])


@pytest.mark.parametrize("brands", [["yandou"], ["烟斗仓库", "yandou"], ["ni"], ["custom"], ["cbanner_womens"], ["烟斗", "custom"], []])
def test_inbound_brand_filter_uses_the_same_record_scope_as_inventory(brand_repository, brands):
    _add_inbound_brand_details(brand_repository)
    expected = brand_repository.list_records(brands=brands, document_types=["进货单", "进货退货单"], page=1, page_size=100)
    result = brand_repository.list_purchase_inbound_details(brands=brands, page=1, page_size=100)
    assert {item["document_id"] for item in result["items"]} == {item["id"] for item in expected["items"]}
    assert result["total"] == len(result["items"]) == expected["total"]


def test_inbound_brand_pagination_totals_and_warehouse_filters_stay_in_sync(brand_repository):
    _add_inbound_brand_details(brand_repository)
    first = brand_repository.list_purchase_inbound_details(brands=["烟斗"], page=1, page_size=1)
    last = brand_repository.list_purchase_inbound_details(brands=["烟斗"], page=3, page_size=1)
    assert first["total"] == last["total"] == 3
    assert first["totals"]["all"] == last["totals"]["all"] == {"purchase_quantity": "0", "purchase_amount": "0", "retail_amount": ""}
    assert last["items"][0]["purchase_quantity"] == "-2"
    assert last["items"][0]["purchase_amount"] == "-20"
    assert brand_repository.list_purchase_inbound_details(brands=["烟斗"], warehouse=["NI总仓"], page=1, page_size=50)["total"] == 0
    narrowed = brand_repository.list_purchase_inbound_details(brands=["烟斗"], document_type="进货退货单", date_start="2026-10-02", product_code="TES", color_name="黑", page=1, page_size=50)
    assert [item["document_id"] for item in narrowed["items"]] == [18]


def test_inbound_export_uses_brand_scope_for_excel_and_operation_log(brand_repository, monkeypatch):
    _add_inbound_brand_details(brand_repository)
    operation_log = Mock()
    monkeypatch.setattr("api.routes.inventory.write_operation_log", operation_log)
    request = SimpleNamespace(app=SimpleNamespace(state=SimpleNamespace(inventory_repository=brand_repository)))
    response = export_purchase_inbound_details(request, brands=["烟斗", "custom"], warehouse=None)

    async def read_content():
        return b"".join([chunk async for chunk in response.body_iterator])

    workbook = load_workbook(BytesIO(asyncio.run(read_content())), data_only=True)
    try:
        rows = list(workbook["商品进货明细"].iter_rows(values_only=True))
        assert {row[rows[0].index("单据编号")] for row in rows[1:]} == {"DOC-1", "DOC-7", "DOC-18", "DOC-19"}
        assert any("品牌 烟斗, custom" == value for value in operation_log.call_args.kwargs["after_data"]["filters"])
    finally:
        workbook.close()


@pytest.mark.parametrize("export", [False, True])
@pytest.mark.parametrize("brands", [[], ["cbanner_mens", "yandou"]])
def test_inbound_routes_forward_repeated_brand_and_warehouse_parameters(monkeypatch, export, brands):
    repository = Mock()
    repository.list_purchase_inbound_details.return_value = {"items": [], "total": 0}
    monkeypatch.setattr("api.routes.inventory.write_operation_log", Mock())
    app = FastAPI()
    app.state.inventory_repository = repository
    path = "/inventory-reports/purchase-inbound-details" + ("/export" if export else "")
    handler = export_purchase_inbound_details if export else list_purchase_inbound_details
    app.add_api_route(path, handler, methods=["GET"])
    params = [("brands", brand) for brand in brands]
    params.extend([("warehouse", "总仓"), ("warehouse", "分仓"), ("supplier", "供应商")])
    if not export:
        params.extend([("page", "2"), ("page_size", "20")])

    with TestClient(app) as client:
        response = client.get(path, params=params)

    assert response.status_code == 200
    repository.list_purchase_inbound_details.assert_called_once()
    filters = repository.list_purchase_inbound_details.call_args.kwargs
    assert filters["brands"] == brands
    assert filters["warehouse"] == ["总仓", "分仓"]
    assert filters["supplier"] == "供应商"
    assert filters["page"] == (1 if export else 2)
    assert filters["page_size"] == (100_000 if export else 20)
