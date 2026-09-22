from __future__ import annotations

import io
from datetime import datetime
from decimal import Decimal
from types import SimpleNamespace
from unittest.mock import Mock

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import Column, DateTime, Integer, MetaData, Numeric, Table, Text, create_engine
from sqlalchemy.pool import StaticPool

from api.auth_middleware import auth_middleware, required_permission_for_request
from api.routes import import_export
from domain import excluded_skus
from storage.auth_repository import DEFAULT_ROLES


@pytest.fixture
def price_export_client(monkeypatch):
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    metadata = MetaData()
    tables = {
        brand: Table(
            brand,
            metadata,
            Column("id", Integer, primary_key=True),
            Column("sku", Text),
            Column("original_sku", Text),
            Column("product_name", Text),
            Column("color", Text),
            Column("supplier_name", Text),
            Column("cost", Numeric(18, 2)),
            Column("factory_sku", Text),
            Column("year", Text),
            Column("created_at", DateTime),
            Column("last_imported_at", DateTime),
            Column("deleted_at", DateTime),
        )
        for brand in ("eblan", "manual_brand")
    }
    metadata.create_all(engine)
    with engine.begin() as connection:
        connection.execute(tables["eblan"].insert(), [
            {
                "id": 1, "sku": "ER741079D24", "original_sku": "ORIGINAL-1",
                "product_name": "女单鞋", "color": "白色", "supplier_name": "168（伊伴女鞋）",
                "cost": Decimal("110.25"), "factory_sku": "00828", "year": "2026",
                "created_at": datetime(2026, 9, 18), "deleted_at": None,
            },
            {
                "id": 2, "sku": "ER741079D35", "original_sku": "ORIGINAL-2",
                "product_name": "女单鞋", "color": "白黑色", "supplier_name": "168（伊伴女鞋）",
                "cost": Decimal("0"), "factory_sku": "828-2", "year": "26年春季款",
                "created_at": datetime(2026, 9, 19), "deleted_at": None,
            },
            {
                "id": 3, "sku": "000123", "original_sku": "OLD-3",
                "product_name": "凉鞋", "color": None, "supplier_name": None,
                "cost": None, "factory_sku": None, "year": "2025",
                "created_at": datetime(2026, 9, 20), "deleted_at": None,
            },
            {
                "id": 4, "sku": "ER-DELETED", "original_sku": "DELETED",
                "product_name": "已删除", "color": "黑色", "supplier_name": "已删除",
                "cost": Decimal("999"), "factory_sku": "DELETED", "year": "2026",
                "created_at": datetime(2026, 9, 20), "deleted_at": datetime(2026, 9, 20),
            },
        ])
        connection.execute(tables["manual_brand"].insert(), {
            "id": 1, "sku": "MANUAL-001", "original_sku": "MANUAL-ORIGINAL",
            "color": "蓝色", "supplier_name": "自定义品牌供应商", "cost": Decimal("88.50"),
            "factory_sku": "00001", "year": "2026",
        })
    repository = SimpleNamespace(
        engine=engine,
        _table_for_brand=tables.__getitem__,
        product_archive_brands=lambda: list(tables),
        is_product_archive_brand=lambda brand: brand in tables,
    )
    user = {"department_code": "商品部", "role_code": "product_user", "permissions": ["product.export"]}
    log = Mock()
    monkeypatch.setattr(import_export, "write_operation_log", log)
    monkeypatch.setattr(excluded_skus, "EXCLUDED_SKUS", frozenset())
    app = FastAPI()
    app.state.repository = repository
    app.state.auth_repository = Mock(
        has_users=Mock(return_value=True),
        get_user_by_session=Mock(return_value=user),
    )
    app.middleware("http")(auth_middleware)
    app.include_router(import_export.router)
    try:
        with TestClient(app) as client:
            yield SimpleNamespace(client=client, user=user, log=log, auth_repository=app.state.auth_repository)
    finally:
        engine.dispose()


@pytest.fixture
def finance_price_export_client(price_export_client):
    role = next(item for item in DEFAULT_ROLES if item["code"] == "finance_user")
    price_export_client.user.update(
        department_code=role["department_code"],
        role_code=role["code"],
        permissions=role["permissions"].split(","),
    )
    return price_export_client


def _export(fixture, **params):
    response = fixture.client.get("/export", params={"brand": "eblan", "mode": "price", **params})
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return load_workbook(io.BytesIO(response.content), data_only=True)


def test_price_export_has_only_five_requested_columns_and_archive_cost(price_export_client):
    workbook = _export(price_export_client)
    try:
        assert workbook.sheetnames == ["物价"]
        assert list(workbook.active.values) == [
            ("货号", "商品全名", "主供应商", "预设售价", "工厂货号"),
            ("000123", "000123", None, None, None),
            ("ER741079D35", "ER741079D35白黑色", "168（伊伴女鞋）", 0, "828-2"),
            ("ER741079D24", "ER741079D24白色", "168（伊伴女鞋）", 110.25, "00828"),
        ]
        assert workbook.active["A2"].number_format == "@"
        assert workbook.active["E4"].number_format == "@"
        assert workbook.active["D4"].data_type == "n"
        assert workbook.active.freeze_panes == "A2"
        assert workbook.active.auto_filter.ref == "A1:E4"
    finally:
        workbook.close()
    price_export_client.log.assert_called_once()
    assert price_export_client.log.call_args.kwargs["after_data"]["exported_rows"] == 3


@pytest.mark.parametrize("params, expected", [
    ({"year": "2026"}, ["ER741079D35", "ER741079D24"]),
    ({"query": " ORIGINAL-1 "}, ["ER741079D24"]),
    ({"query": "D35\n000123"}, ["000123", "ER741079D35"]),
    ({"sku_prefix": "ER", "query": "D35", "year": "2026"}, ["ER741079D35"]),
    ({"ids": "1,4", "query": "not-found", "year": "2025"}, ["ER741079D24"]),
    ({"sku_prefix": "not-found"}, []),
])
def test_price_export_respects_filters_selection_and_deleted_records(price_export_client, params, expected):
    workbook = _export(price_export_client, **params)
    try:
        assert [row[0] for row in list(workbook.active.values)[1:]] == expected
        assert workbook.active.max_column == 5
    finally:
        workbook.close()


def test_price_export_all_brands_includes_custom_brand(price_export_client):
    workbook = _export(price_export_client, brand="all", year="2026")
    try:
        rows = list(workbook.active.values)
        assert [row[0] for row in rows[1:]] == ["ER741079D35", "ER741079D24", "MANUAL-001"]
        assert rows[-1] == ("MANUAL-001", "MANUAL-001蓝色", "自定义品牌供应商", 88.5, "00001")
        assert workbook.active.max_column == 5
    finally:
        workbook.close()


@pytest.mark.parametrize("method", ["get", "head"])
@pytest.mark.parametrize("department", ["客服部", "美工部"])
@pytest.mark.parametrize("permission", ["product.export", "product.price_export"])
def test_price_export_denies_cost_restricted_departments(price_export_client, department, method, permission):
    price_export_client.user["department_code"] = department
    price_export_client.user["permissions"] = [permission]
    response = getattr(price_export_client.client, method)("/export", params={"brand": "eblan", "mode": "price"})
    assert response.status_code == 403
    price_export_client.log.assert_not_called()


def test_price_export_super_admin_can_export_even_in_restricted_department(price_export_client):
    price_export_client.user.update({"department_code": "美工部", "role_code": "super_admin"})
    workbook = _export(price_export_client)
    assert workbook.active["D4"].value == 110.25
    workbook.close()


def test_price_export_preflight_succeeds_without_generating_workbook(price_export_client):
    response = price_export_client.client.head("/export", params={"brand": "eblan", "mode": "price"})
    assert response.status_code == 200
    assert response.content == b""
    price_export_client.log.assert_not_called()


@pytest.mark.parametrize("params", [{"ids": "oops"}, {"brand": "unknown"}])
def test_price_export_rejects_invalid_parameters(price_export_client, params):
    response = price_export_client.client.get("/export", params={"brand": "eblan", "mode": "price", **params})
    assert response.status_code == 400
    price_export_client.log.assert_not_called()


def test_price_export_reuses_existing_export_permission():
    assert required_permission_for_request("GET", "/export") == "product.export"
    assert required_permission_for_request("HEAD", "/export") == "product.export"


def test_finance_can_export_selected_prices_with_archive_cost(finance_price_export_client):
    workbook = _export(finance_price_export_client, ids="1")
    try:
        assert list(workbook.active.values) == [
            ("货号", "商品全名", "主供应商", "预设售价", "工厂货号"),
            ("ER741079D24", "ER741079D24白色", "168（伊伴女鞋）", 110.25, "00828"),
        ]
    finally:
        workbook.close()


def test_finance_price_export_preflight_succeeds(finance_price_export_client):
    response = finance_price_export_client.client.head("/export?brand=eblan&mode=price")
    assert response.status_code == 200
    assert response.content == b""
    finance_price_export_client.log.assert_not_called()


@pytest.mark.parametrize("method", ["get", "head"])
@pytest.mark.parametrize("query", [
    "", "&mode=", "&mode=with_sizes", "&mode=unknown", "&mode=PRICE", "&mode=price%20",
    "&mode=price&mode=with_sizes", "&mode=price&mode=",
])
def test_finance_cannot_bypass_price_only_export_permission(finance_price_export_client, method, query):
    response = getattr(finance_price_export_client.client, method)(f"/export?brand=eblan{query}")
    assert response.status_code == 403
    finance_price_export_client.log.assert_not_called()


@pytest.mark.parametrize("method", ["get", "head"])
def test_price_export_without_session_requires_login(price_export_client, method):
    price_export_client.auth_repository.get_user_by_session.return_value = None
    response = getattr(price_export_client.client, method)("/export?brand=eblan&mode=price")
    assert response.status_code == 401
    price_export_client.log.assert_not_called()


@pytest.mark.parametrize("method", ["get", "head"])
def test_product_view_permission_alone_cannot_export_prices(finance_price_export_client, method):
    finance_price_export_client.user["permissions"] = ["product.view"]
    response = getattr(finance_price_export_client.client, method)("/export?brand=eblan&mode=price")
    assert response.status_code == 403
    finance_price_export_client.log.assert_not_called()
