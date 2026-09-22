from __future__ import annotations

import asyncio
import io
from decimal import Decimal
from types import SimpleNamespace
from unittest.mock import Mock

import pytest
from openpyxl import load_workbook

from api.routes.inventory import export_inventory
from domain.inventory_sources import ACCOUNTING_DOCUMENT_TYPES
from storage.inventory_repository import InventoryRepository


def _record(record_id=1, document_type="应付款减少", amount="1000"):
    return {
        "id": record_id,
        "document_number": f"TEST-{record_id:04d}",
        "date": "2026-09-03",
        "supplier": "测试供应商（千百度女鞋）",
        "document_type": document_type,
        "amount": Decimal(amount),
        "summary": "质量罚款",
        "handler": "测试经手人",
        "additional_note": "需工厂确认",
    }


def _detail(document_id=1, amount="1000"):
    return {
        "id": document_id,
        "document_id": document_id,
        "product_name": "质量罚款科目",
        "amount": Decimal(amount),
        "remark": "鞋面瑕疵扣款",
    }


def _repository(records, details):
    repository = Mock(spec=InventoryRepository)
    repository.list_records.return_value = {"items": records, "total": len(records)}
    repository.list_details_for_documents.side_effect = lambda document_ids: [
        detail for detail in details if detail["document_id"] in document_ids
    ]
    return repository


def _export(repository, **filters):
    request = SimpleNamespace(app=SimpleNamespace(state=SimpleNamespace(inventory_repository=repository)))
    response = export_inventory(request, **filters)

    async def read_content():
        return b"".join([chunk async for chunk in response.body_iterator])

    assert response.media_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return load_workbook(io.BytesIO(asyncio.run(read_content())), data_only=True)


def _rows(worksheet):
    values = list(worksheet.iter_rows(values_only=True))
    return [dict(zip(values[0], row)) for row in values[1:]]


def test_export_all_types_with_fine_filter_keeps_all_accounting_records():
    records = [_record(record_id) for record_id in range(1, 31)]
    repository = _repository(records, [_detail(record_id) for record_id in range(1, 31)])

    workbook = _export(repository, summary="罚款", completion_status="completed", exclude_document_type="进货订单")
    try:
        summary_rows = _rows(workbook["经营历程"])
        detail_rows = _rows(workbook["单据明细"])
        assert len(summary_rows) == 30
        assert len(detail_rows) == 30
        assert summary_rows[0]["单据编号"] == "TEST-0001"
        assert summary_rows[0]["单位全名"] == records[0]["supplier"]
        assert summary_rows[0]["金额"] == 1000
        assert summary_rows[0]["摘要"] == "质量罚款"
        assert summary_rows[0]["附加说明"] == "需工厂确认"
        assert detail_rows[0]["费用项目名 / 科目"] == "质量罚款科目"
        assert detail_rows[0]["金额"] == 1000
        assert detail_rows[0]["备注"] == "鞋面瑕疵扣款"
        assert "货号" not in detail_rows[0]
        assert "单价" not in detail_rows[0]
        assert "交货日期" not in detail_rows[0]
    finally:
        workbook.close()
    repository.list_details_for_documents.assert_called_once_with(list(range(1, 31)))
    filters = repository.list_records.call_args.kwargs
    assert filters["summary"] == "罚款"
    assert filters["completion_status"] == "completed"
    assert filters["exclude_document_type"] == "进货订单"
    assert filters["page"] == 1
    assert filters["page_size"] == 100_000


@pytest.mark.parametrize("document_type", ACCOUNTING_DOCUMENT_TYPES)
def test_export_specific_accounting_type_includes_summary_and_details(document_type):
    repository = _repository([_record(document_type=document_type, amount="123.45")], [_detail(amount="123.45")])

    workbook = _export(repository, document_type=document_type)
    try:
        assert _rows(workbook["经营历程"])[0]["单据类型"] == document_type
        detail = _rows(workbook["单据明细"])[0]
        assert detail["单据类型"] == document_type
        assert detail["日期"] == "2026-09-03"
        assert detail["金额"] == 123.45
        assert detail["备注"] == "鞋面瑕疵扣款"
    finally:
        workbook.close()
    assert repository.list_records.call_args.kwargs["document_type"] == document_type


def test_export_selected_accounting_records_limits_both_sheets():
    repository = _repository([_record(1), _record(2)], [_detail(1), _detail(2)])

    workbook = _export(repository, ids="2", summary="罚款", completion_status="completed")
    try:
        assert [row["单据编号"] for row in _rows(workbook["经营历程"])] == ["TEST-0002"]
        assert [row["单据编号"] for row in _rows(workbook["单据明细"])] == ["TEST-0002"]
    finally:
        workbook.close()
    repository.list_details_for_documents.assert_called_once_with([2])


def test_export_mixed_accounting_and_product_records_preserves_both_details():
    product_detail = {
        **_detail(2, "30"),
        "product_code": "SHOE-001",
        "product_name": "测试鞋",
        "quantity": "3",
        "unit_price": "10",
        "size_quantities": {"230": "3"},
        "remark": "商品明细备注",
    }
    repository = _repository([_record(1), _record(2, "进货单", "30")], [_detail(1), product_detail])

    workbook = _export(repository)
    try:
        assert len(_rows(workbook["经营历程"])) == 2
        accounting, product = _rows(workbook["单据明细"])
        assert accounting["商品全名 / 费用项目名 / 科目"] == "质量罚款科目"
        assert accounting["金额"] == 1000
        assert accounting["备注"] == "鞋面瑕疵扣款"
        assert product["货号"] == "SHOE-001"
        assert product["230"] == "3"
        assert product["数量"] == "3"
        assert product["金额"] == 30
        assert product["备注"] == "商品明细备注"
    finally:
        workbook.close()


def test_export_accounting_record_without_details_keeps_summary():
    repository = _repository([_record()], [])

    workbook = _export(repository)
    try:
        assert len(_rows(workbook["经营历程"])) == 1
        assert _rows(workbook["经营历程"])[0]["金额"] == 1000
        assert _rows(workbook["单据明细"]) == []
    finally:
        workbook.close()


@pytest.mark.parametrize("amount", ["0", "-123.45"])
def test_export_accounting_amount_preserves_zero_and_negative_values(amount):
    repository = _repository([_record(amount=amount)], [_detail(amount=amount)])

    workbook = _export(repository, completion_status="incomplete")
    try:
        assert _rows(workbook["经营历程"])[0]["金额"] == float(amount)
        assert _rows(workbook["单据明细"])[0]["金额"] == float(amount)
    finally:
        workbook.close()
    assert repository.list_records.call_args.kwargs["completion_status"] == "incomplete"


def test_export_product_only_records_keeps_existing_columns():
    repository = _repository([_record(document_type="进货单")], [{
        **_detail(), "product_code": "SHOE-001", "product_name": "测试鞋",
    }])

    workbook = _export(repository, document_type="进货单")
    try:
        assert _rows(workbook["经营历程"])[0]["供应商"] == _record()["supplier"]
        detail = _rows(workbook["单据明细"])[0]
        assert detail["商品全名"] == "测试鞋"
        assert detail["货号"] == "SHOE-001"
        assert "订货日期" in detail
        assert "交货日期" in detail
    finally:
        workbook.close()
