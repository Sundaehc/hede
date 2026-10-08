from io import BytesIO
from datetime import date
import asyncio
import json
from types import SimpleNamespace
from unittest.mock import AsyncMock, Mock

import pytest
from fastapi import HTTPException, UploadFile
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine, event, select

from domain.inventory_template_import import TEMPLATE_HEADERS, TemplateDocument, read_template_documents
from domain.inventory_schema import GENERAL_CUSTOMER_SHOP_TABLE, GENERAL_CUSTOMER_UNIT_TABLE, INVENTORY_ACCOUNT_SUBJECT_TABLE, INVENTORY_DETAIL_TABLE, INVENTORY_TABLE, SUPPLIER_TABLE, WAREHOUSE_BRAND_TABLE, WAREHOUSE_TABLE
from api.routes import inventory as inventory_routes
from storage.inventory_repository import InventoryRepository


def _workbook(headers, rows):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "导出数据"
    worksheet.append(headers)
    for row in rows:
        worksheet.append(row)
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _seed_template_purchase_master_data(engine):
    SUPPLIER_TABLE.create(engine)
    WAREHOUSE_TABLE.create(engine)
    with engine.begin() as connection:
        connection.execute(SUPPLIER_TABLE.insert(), [
            {"id": 1, "brand": "ni", "name": "供应商A"},
            {"id": 2, "brand": "ni", "name": "供应商B"},
        ])
        connection.execute(WAREHOUSE_TABLE.insert(), [
            {"id": 1, "name": "仓库A"},
            {"id": 2, "name": "仓库B"},
        ])


def test_all_downloadable_templates_omit_document_number_and_system_code():
    assert len(TEMPLATE_HEADERS) == 7
    assert all("单据编号" not in headers and "系统码" not in headers for headers in TEMPLATE_HEADERS.values())
    assert "商品编码" in TEMPLATE_HEADERS["sale"]
    assert "商品编码" in TEMPLATE_HEADERS["sale_return"]
    assert {headers[0] for headers in TEMPLATE_HEADERS.values()} == {"日期"}
    assert {headers[2] for headers in TEMPLATE_HEADERS.values()} == {"单位全名", "经手人", "仓库全名"}
    assert TEMPLATE_HEADERS["stock_loss"] == TEMPLATE_HEADERS["stock_gain"] == (
        "日期", "单据类型", "仓库全名", "经手人", "摘要", "商品编码", "数量",
    )
    assert all("数量" in headers for headers in TEMPLATE_HEADERS.values() if headers[-1] == "单价")


@pytest.mark.parametrize("kind, document_type", [
    ("stock_loss", "报损单"),
    ("stock_gain", "报溢单"),
])
def test_stock_adjustment_template_download_and_read(kind, document_type):
    response = inventory_routes.download_inventory_template(kind)

    async def read_content():
        return b"".join([chunk async for chunk in response.body_iterator])

    workbook = load_workbook(BytesIO(asyncio.run(read_content())), read_only=True)
    try:
        assert tuple(cell.value for cell in workbook.active[1]) == TEMPLATE_HEADERS[kind]
    finally:
        workbook.close()

    content = _workbook(TEMPLATE_HEADERS[kind], [
        [date(2025, 4, 30), document_type, "NI仙岩仓库", "陈章瑜", "月末调整", "NI24Q1A02020139", 2],
        [date(2025, 4, 30), document_type, "NI仙岩仓库", "陈章瑜", "月末调整", "NI24Q1A02020436", 1],
    ])
    parsed_kind, _, documents = read_template_documents(content)
    assert parsed_kind == kind
    assert len(documents) == 1
    assert (documents[0].date, documents[0].document_type, documents[0].supplier) == (
        "2025-04-30", document_type, "",
    )
    assert documents[0].warehouse == "NI仙岩仓库"
    assert documents[0].rows == [
        {"product_code": "NI24Q1A02020139", "quantity": "2", "unit_price": ""},
        {"product_code": "NI24Q1A02020436", "quantity": "1", "unit_price": ""},
    ]


@pytest.mark.parametrize("kind, document_type", [
    ("stock_loss", "报损单"),
    ("stock_gain", "报溢单"),
])
def test_stock_adjustment_template_rejects_wrong_document_type(kind, document_type):
    wrong_type = "报溢单" if document_type == "报损单" else "报损单"
    content = _workbook(TEMPLATE_HEADERS[kind], [
        ["2025-04-30", document_type, "NI仙岩仓库", "陈章瑜", "月末调整", "NI24Q1A02020139", 2],
        ["2025-04-30", wrong_type, "NI仙岩仓库", "陈章瑜", "月末调整", "NI24Q1A02020436", 1],
    ])
    with pytest.raises(HTTPException, match="单据类型与模板不匹配"):
        read_template_documents(content)


@pytest.mark.parametrize("kind, document_type", [
    ("stock_loss", "报损单"),
    ("stock_gain", "报溢单"),
])
def test_stock_adjustment_template_preview_needs_warehouse_but_no_supplier(kind, document_type):
    content = _workbook(TEMPLATE_HEADERS[kind], [
        ["2025-04-30", document_type, "NI仙岩仓库", "陈章瑜", "月末调整", "NI24Q1A02020139", 2],
    ])
    repository = Mock(spec=InventoryRepository)
    repository.list_warehouses.return_value = [{"name": "NI仙岩仓库"}]
    repository.get_warehouse_by_name.return_value = {"name": "NI仙岩仓库", "brand": "ni"}
    repository.preview_template_documents.return_value = {"total": 1, "conflicts": []}
    request = SimpleNamespace(app=SimpleNamespace(state=SimpleNamespace(inventory_repository=repository)))

    preview = asyncio.run(inventory_routes.preview_inventory_template(
        request, UploadFile(filename="调整.xlsx", file=BytesIO(content)),
    ))
    assert preview["kind"] == kind
    assert preview["brand_by_key"]
    repository.list_suppliers.assert_not_called()

    repository.list_warehouses.return_value = []
    with pytest.raises(HTTPException, match="仓库"):
        asyncio.run(inventory_routes.preview_inventory_template(
            request, UploadFile(filename="调整.xlsx", file=BytesIO(content)),
        ))


@pytest.mark.parametrize("kind, document_type", [
    ("stock_loss", "报损单"),
    ("stock_gain", "报溢单"),
])
def test_stock_adjustment_template_import_creates_document_without_supplier(kind, document_type, monkeypatch):
    content = _workbook(TEMPLATE_HEADERS[kind], [
        ["2025-04-30", document_type, "NI仙岩仓库", "陈章瑜", "月末调整", "NI24Q1A02020139", 2],
    ])
    engine = create_engine("sqlite://")
    event.listen(engine, "connect", lambda connection, _record: connection.create_function(
        "date_trunc", 2, lambda _unit, value: value,
    ))
    INVENTORY_TABLE.create(engine)
    INVENTORY_DETAIL_TABLE.create(engine)
    WAREHOUSE_BRAND_TABLE.create(engine)
    WAREHOUSE_TABLE.create(engine)
    with engine.begin() as connection:
        connection.execute(WAREHOUSE_TABLE.insert().values(id=1, name="NI仙岩仓库", brand="ni"))
    repository = object.__new__(InventoryRepository)
    repository.engine = engine
    repository._prepare_record = lambda record: {**InventoryRepository._prepare_record(record), "id": 1}
    build_details = Mock(return_value=[
        {"id": 1, "product_code": "NI24Q1A020201", "quantity": 2, "unit_price": 10, "amount": 20},
    ])
    monkeypatch.setattr(inventory_routes, "_build_purchase_details_from_rows", build_details)
    monkeypatch.setattr(inventory_routes, "write_operation_log", lambda *_args, **_kwargs: None)
    request = SimpleNamespace(
        app=SimpleNamespace(state=SimpleNamespace(inventory_repository=repository)),
        form=AsyncMock(return_value={"decisions": "{}"}),
    )
    try:
        result = asyncio.run(inventory_routes.import_inventory_template(
            request, UploadFile(filename="调整.xlsx", file=BytesIO(content)),
        ))
        with engine.connect() as connection:
            record = connection.execute(select(INVENTORY_TABLE)).mappings().one()
            detail = connection.execute(select(INVENTORY_DETAIL_TABLE)).mappings().one()
        assert (result["created"], result["details"]) == (1, 1)
        assert record["document_type"] == document_type
        assert record["supplier"] is None
        assert record["warehouse"] == "NI仙岩仓库"
        assert record["raw_payload"]["brand"] == "ni"
        assert record["total_count"] == 2
        assert record["amount"] == 20
        assert detail["product_code"] == "NI24Q1A020201"
        assert build_details.call_args.args[1] == [
            {"product_code": "NI24Q1A02020139", "quantity": "2", "unit_price": ""},
        ]
        assert build_details.call_args.kwargs["brand"] == "ni"
    finally:
        engine.dispose()


def test_read_legacy_template_headers_are_still_supported():
    content = _workbook(
        ("单据日期", "单据类型", "单位全名", "仓库全名", "制单人", "摘要", "商品编码", "数量", "单价"),
        [["2026-10-01", "进货单", "供应商A", "仓库A", "张三", "进货一批", "A001", 1, 10]],
    )

    kind, _, documents = read_template_documents(content)

    assert kind == "purchase"
    assert documents[0].date == "2026-10-01"


def test_read_template_rejects_mixed_document_types():
    content = _workbook(
        TEMPLATE_HEADERS["sale"],
        [
            ["2026-10-01", "批发销售单", "客户A", "仓库A", "张三", "导入", "A001", 1, 10],
            ["2026-10-01", "进货单", "供应商A", "仓库A", "张三", "导入", "A002", 1, 10],
        ],
    )

    with pytest.raises(HTTPException, match="单据类型与模板不匹配"):
        read_template_documents(content)


def test_read_sale_template_groups_details_without_number():
    content = _workbook(
        TEMPLATE_HEADERS["sale"],
        [
            ["2026-10-01", "批发销售单", "品牌A客户", "品牌A仓库", "张三", "销售一批", "A001", 2, 99],
            ["2026-10-01", "批发销售单", "品牌A客户", "品牌A仓库", "张三", "销售一批", "A002", 1, 100],
        ],
    )

    kind, sheet, documents = read_template_documents(content)

    assert kind == "sale"
    assert sheet == "导出数据"
    assert len(documents) == 1
    assert documents[0].summary == "销售一批"
    assert documents[0].key
    assert len(documents[0].rows) == 2


def test_read_accounting_template_has_no_product_code_requirement():
    content = _workbook(
        TEMPLATE_HEADERS["accounting"],
        [["2026/10/01", "应付款增加", "张三", "供应商A", "包装费", "包装费用", 1200]],
    )

    kind, _, documents = read_template_documents(content)

    assert kind == "accounting"
    assert documents[0].date == "2026-10-01"
    assert documents[0].rows == [{"product_name": "包装费用", "amount": "1200"}]


def test_read_template_separates_different_summaries():
    content = _workbook(
        TEMPLATE_HEADERS["sale"],
        [
            ["2026-10-01", "批发销售单", "客户A", "仓库A", "张三", "销售一批", "A001", 1, 99],
            ["2026-10-01", "批发销售单", "客户A", "仓库A", "张三", "销售二批", "A002", 1, 99],
        ],
    )

    _, _, documents = read_template_documents(content)

    assert len(documents) == 2
    assert documents[0].key != documents[1].key


def test_read_legacy_number_column_is_ignored():
    content = _workbook(
        (*("系统码" if header == "商品编码" else header for header in TEMPLATE_HEADERS["sale"]), "单据编号"),
        [
            ["2026-10-01", "批发销售单", "客户A", "仓库A", "张三", "销售一批", "A001", 1, 99, "SOURCE-001"],
            ["2026-10-01", "批发销售单", "客户A", "仓库A", "张三", "销售一批", "A002", 1, 99, "SOURCE-002"],
        ],
    )

    _, _, documents = read_template_documents(content)

    assert len(documents) == 1
    assert len(documents[0].rows) == 2
    assert "SOURCE-001" not in documents[0].key
    assert documents[0].rows[0]["product_code"] == "A001"


def test_read_legacy_sale_return_system_code_without_number():
    headers = tuple("系统码" if header == "商品编码" else header for header in TEMPLATE_HEADERS["sale_return"])
    content = _workbook(
        headers,
        [["2026-10-01", "批发销售退货单", "客户A", "仓库A", "张三", "销售退货", "A001", 1, 99]],
    )

    kind, _, documents = read_template_documents(content)

    assert kind == "sale_return"
    assert documents[0].rows[0]["product_code"] == "A001"


def test_template_overwrite_replaces_details_and_keeps_document_number():
    engine = create_engine("sqlite://")
    event.listen(engine, "connect", lambda connection, _record: connection.create_function(
        "date_trunc", 2, lambda _unit, value: value,
    ))
    INVENTORY_TABLE.create(engine)
    INVENTORY_DETAIL_TABLE.create(engine)
    _seed_template_purchase_master_data(engine)
    repository = object.__new__(InventoryRepository)
    repository.engine = engine
    with engine.begin() as connection:
        connection.execute(INVENTORY_TABLE.insert().values(
            id=1, date="2026-10-01", date_value=date(2026, 10, 1),
            document_number="JHD-2026-10-01-0001", document_type="进货单",
            supplier="供应商A", warehouse="仓库A", handler="原经手人", summary="进货一批",
        ))
        connection.execute(INVENTORY_DETAIL_TABLE.insert().values(
            id=1, document_id=1, product_code="OLD", quantity=3, amount=30,
        ))
    plan = {
        "date": "2026-10-01", "document_type": "进货单", "supplier": "供应商A",
        "warehouse": "仓库A", "handler": "新经手人", "summary": "进货一批",
        "source_workbook": "导入.xlsx", "source_sheet": "导出数据", "brand": "ni",
        "decision": "overwrite", "details": [{"id": 2, "product_code": "NEW", "quantity": 2, "amount": 20}],
    }
    try:
        result = repository.import_template_documents([plan])
        with engine.connect() as connection:
            record = connection.execute(select(INVENTORY_TABLE)).mappings().one()
            details = connection.execute(select(INVENTORY_DETAIL_TABLE)).mappings().all()
        assert result == {"created": 0, "replaced": 1, "skipped": 0, "details": 1}
        assert record["document_number"] == "JHD-2026-10-01-0001"
        assert record["warehouse"] == "仓库A"
        assert record["handler"] == "新经手人"
        assert record["total_count"] == 2
        assert record["amount"] == 20
        assert [detail["product_code"] for detail in details] == ["NEW"]

        with pytest.raises(ValueError, match="请先确认覆盖或新增"):
            repository.import_template_documents([{**plan, "decision": None}])
        with engine.connect() as connection:
            assert [row[0] for row in connection.execute(select(INVENTORY_DETAIL_TABLE.c.product_code))] == ["NEW"]
    finally:
        engine.dispose()


def test_regular_import_overwrite_replaces_details_and_keeps_document_number():
    engine = create_engine("sqlite://")
    event.listen(engine, "connect", lambda connection, _record: connection.create_function(
        "date_trunc", 2, lambda _unit, value: value,
    ))
    INVENTORY_TABLE.create(engine)
    INVENTORY_DETAIL_TABLE.create(engine)
    repository = object.__new__(InventoryRepository)
    repository.engine = engine
    with engine.begin() as connection:
        connection.execute(INVENTORY_TABLE.insert().values(
            id=1, date="2026-10-01", date_value=date(2026, 10, 1),
            document_number="JHD-2026-10-01-0001", document_type="进货单",
            supplier="供应商A", warehouse="仓库A", handler="原经手人", summary="进货一批",
        ))
        connection.execute(INVENTORY_DETAIL_TABLE.insert().values(
            id=1, document_id=1, product_code="OLD", quantity=3, amount=30,
        ))
    try:
        updated = repository.overwrite_imported_document(
            1, expected_date="2026-10-01", expected_summary="进货一批",
            expected_document_type="进货单", expected_supplier="供应商A", expected_warehouse="仓库A",
            record={"date": "2026-10-01", "document_type": "进货单", "supplier": "供应商A",
                    "warehouse": "仓库A", "handler": "新经手人", "summary": "进货一批",
                    "source_workbook": "导入.xlsx", "source_sheet": "导出数据", "source_row_number": "import_purchase"},
            details=[{"id": 2, "product_code": "NEW", "quantity": 2, "amount": 20}],
        )
        assert updated is not None
        assert updated["document_number"] == "JHD-2026-10-01-0001"
        assert updated["warehouse"] == "仓库A"
        assert updated["handler"] == "新经手人"
        assert updated["total_count"] == 2
        assert updated["amount"] == 20
        with engine.connect() as connection:
            assert [row[0] for row in connection.execute(select(INVENTORY_DETAIL_TABLE.c.product_code))] == ["NEW"]

        rejected = repository.overwrite_imported_document(
            1, expected_date="2026-10-01", expected_summary="错误摘要",
            expected_document_type="进货单", expected_supplier="供应商A", expected_warehouse="仓库A",
            record={}, details=[],
        )
        assert rejected is None
        with engine.connect() as connection:
            assert [row[0] for row in connection.execute(select(INVENTORY_DETAIL_TABLE.c.product_code))] == ["NEW"]
    finally:
        engine.dispose()


@pytest.mark.parametrize("changed_field, changed_value", [
    ("date", "2026-10-02"),
    ("summary", "进货二批"),
    ("document_type", "进货退货单"),
    ("supplier", "供应商B"),
    ("warehouse", "仓库B"),
])
def test_import_conflict_requires_five_matching_fields(changed_field, changed_value):
    engine = create_engine("sqlite://")
    event.listen(engine, "connect", lambda connection, _record: connection.create_function(
        "date_trunc", 2, lambda _unit, value: value,
    ))
    INVENTORY_TABLE.create(engine)
    INVENTORY_DETAIL_TABLE.create(engine)
    _seed_template_purchase_master_data(engine)
    repository = object.__new__(InventoryRepository)
    repository.engine = engine
    with engine.begin() as connection:
        connection.execute(INVENTORY_TABLE.insert().values(
            id=1, date="2026-10-01", date_value=date(2026, 10, 1),
            document_number="JHD-2026-10-01-0001", document_type="进货单",
            supplier="供应商A", warehouse="仓库A", summary="进货一批",
        ))
    fields = {
        "date": "2026-10-01", "summary": "进货一批", "document_type": "进货单",
        "supplier": "供应商A", "warehouse": "仓库A",
    }
    fields[changed_field] = changed_value
    document = TemplateDocument("test", fields["date"], fields["document_type"],
                                fields["supplier"], fields["warehouse"], "新经手人", fields["summary"], [])
    plan = {
        **fields, "handler": "新经手人", "source_workbook": "导入.xlsx",
        "source_sheet": "导出数据", "brand": "ni", "decision": None,
        "details": [{"id": 2, "product_code": "NEW", "quantity": 2, "amount": 20}],
    }
    try:
        assert repository.preview_template_documents([document])["conflicts"] == []
        assert repository.get_record_for_append(
            date_value=fields["date"], summary=fields["summary"], document_type=fields["document_type"],
            supplier=fields["supplier"], warehouse=fields["warehouse"],
        ) is None
        repository._prepare_record = lambda record: {**InventoryRepository._prepare_record(record), "id": 2}
        assert repository.import_template_documents([plan]) == {"created": 1, "replaced": 0, "skipped": 0, "details": 1}
        with engine.connect() as connection:
            assert len(connection.execute(select(INVENTORY_TABLE.c.id)).all()) == 2
    finally:
        engine.dispose()


def test_template_preview_ignores_handler_but_matches_warehouse():
    engine = create_engine("sqlite://")
    event.listen(engine, "connect", lambda connection, _record: connection.create_function(
        "date_trunc", 2, lambda _unit, value: value,
    ))
    INVENTORY_TABLE.create(engine)
    repository = object.__new__(InventoryRepository)
    repository.engine = engine
    with engine.begin() as connection:
        connection.execute(INVENTORY_TABLE.insert().values(
            id=1, date="2026-10-01", date_value=date(2026, 10, 1),
            document_number="JHD-2026-10-01-0001", document_type="进货单",
            supplier="供应商A", warehouse="仓库A", handler="旧经手人", summary="进货一批",
        ))
    try:
        document = TemplateDocument("test", "2026-10-01", "进货单", "供应商A", "仓库A", "新经手人", "进货一批", [])
        conflicts = repository.preview_template_documents([document])["conflicts"]
        assert len(conflicts) == 1
        assert conflicts[0]["existing_number"] == "JHD-2026-10-01-0001"
        matched = repository.get_record_for_append(
            date_value=document.date, summary=document.summary,
            document_type=document.document_type, supplier=document.supplier,
            warehouse=document.warehouse,
        )
        assert matched is not None
        assert matched["document_number"] == "JHD-2026-10-01-0001"
    finally:
        engine.dispose()


def test_template_preview_same_file_requires_five_matching_fields():
    engine = create_engine("sqlite://")
    event.listen(engine, "connect", lambda connection, _record: connection.create_function(
        "date_trunc", 2, lambda _unit, value: value,
    ))
    INVENTORY_TABLE.create(engine)
    repository = object.__new__(InventoryRepository)
    repository.engine = engine
    base = TemplateDocument("first", "2026-10-01", "进货单", "供应商A", "仓库A", "甲", "进货一批", [])
    try:
        for field, value in (
            ("date", "2026-10-02"), ("summary", "进货二批"),
            ("document_type", "进货退货单"), ("supplier", "供应商B"),
            ("warehouse", "仓库B"),
        ):
            other = TemplateDocument("second", base.date, base.document_type, base.supplier,
                                     base.warehouse, base.handler, base.summary, [])
            setattr(other, field, value)
            assert repository.preview_template_documents([base, other])["conflicts"] == []

        same_fields = TemplateDocument("second", base.date, base.document_type, base.supplier,
                                       base.warehouse, "乙", base.summary, [])
        conflicts = repository.preview_template_documents([base, same_fields])["conflicts"]
        assert len(conflicts) == 1
        assert conflicts[0]["key"] == "second"
        different_warehouse = TemplateDocument("second", base.date, base.document_type, base.supplier,
                                               "仓库B", "乙", base.summary, [])
        conflicts = repository.preview_template_documents([base, different_warehouse])["conflicts"]
        assert conflicts == []
    finally:
        engine.dispose()


def test_template_import_skips_only_selected_document():
    engine = create_engine("sqlite://")
    event.listen(engine, "connect", lambda connection, _record: connection.create_function(
        "date_trunc", 2, lambda _unit, value: value,
    ))
    INVENTORY_TABLE.create(engine)
    INVENTORY_DETAIL_TABLE.create(engine)
    _seed_template_purchase_master_data(engine)
    repository = object.__new__(InventoryRepository)
    repository.engine = engine
    with engine.begin() as connection:
        connection.execute(INVENTORY_TABLE.insert().values(
            id=1, date="2026-10-01", date_value=date(2026, 10, 1),
            document_number="JHD-2026-10-01-0001", document_type="进货单",
            supplier="供应商A", warehouse="仓库A", summary="保留单据",
        ))
        connection.execute(INVENTORY_DETAIL_TABLE.insert().values(
            id=1, document_id=1, product_code="OLD", quantity=3, amount=30,
        ))
    repository._prepare_record = lambda record: {**InventoryRepository._prepare_record(record), "id": 2}
    try:
        result = repository.import_template_documents([
            {"decision": "skip"},
            {
                "date": "2026-10-01", "document_type": "进货单", "supplier": "供应商A",
                "warehouse": "仓库B", "handler": "张三", "summary": "新单据",
                "source_workbook": "导入.xlsx", "source_sheet": "导出数据", "brand": "ni",
                "decision": None, "details": [{"id": 2, "product_code": "NEW", "quantity": 2, "amount": 20}],
            },
        ])
        assert result == {"created": 1, "replaced": 0, "skipped": 1, "details": 1}
        with engine.connect() as connection:
            records = connection.execute(select(INVENTORY_TABLE).order_by(INVENTORY_TABLE.c.id)).mappings().all()
            details = connection.execute(select(INVENTORY_DETAIL_TABLE).order_by(INVENTORY_DETAIL_TABLE.c.id)).mappings().all()
        assert [record["summary"] for record in records] == ["保留单据", "新单据"]
        assert [(detail["document_id"], detail["product_code"]) for detail in details] == [(1, "OLD"), (2, "NEW")]
        assert repository.import_template_documents([{"decision": "skip"}]) == {
            "created": 0, "replaced": 0, "skipped": 1, "details": 0,
        }
    finally:
        engine.dispose()


def test_template_import_cancel_all_does_not_write(monkeypatch):
    content = _workbook(TEMPLATE_HEADERS["accounting"], [
        ["2026-10-01", "应付款增加", "张三", "供应商A", "包装费", "包装费", 100],
    ])
    _, _, documents = read_template_documents(content)
    repository = Mock(spec=InventoryRepository)
    repository.preview_template_documents.return_value = {"conflicts": [{"key": documents[0].key}]}
    request = SimpleNamespace(
        app=SimpleNamespace(state=SimpleNamespace(inventory_repository=repository)),
        form=AsyncMock(return_value={"decisions": json.dumps({documents[0].key: {"action": "skip"}})}),
    )
    monkeypatch.setattr(inventory_routes, "write_operation_log", lambda *_args, **_kwargs: None)
    result = asyncio.run(inventory_routes.import_inventory_template(
        request, UploadFile(filename="应付.xlsx", file=BytesIO(content)),
    ))
    assert (result["created"], result["replaced"], result["skipped"]) == (0, 0, 1)
    repository.import_template_documents.assert_not_called()


def test_template_import_skips_missing_master_data_for_cancelled_document(monkeypatch):
    content = _workbook(TEMPLATE_HEADERS["accounting"], [
        ["2026-10-01", "应付款增加", "张三", "未建供应商", "取消", "未建科目", 100],
        ["2026-10-02", "应付款增加", "张三", "供应商A", "保留", "包装费", 200],
    ])
    _, _, documents = read_template_documents(content)
    repository = Mock(spec=InventoryRepository)
    repository.list_suppliers.return_value = [{"name": "供应商A"}]
    repository.list_account_subjects.return_value = [{"name": "包装费"}]
    repository.preview_template_documents.return_value = {"conflicts": [{"key": documents[0].key}]}
    repository.import_template_documents.return_value = {"created": 1, "replaced": 0, "skipped": 1, "details": 1}
    request = SimpleNamespace(
        app=SimpleNamespace(state=SimpleNamespace(inventory_repository=repository)),
        form=AsyncMock(return_value={"decisions": json.dumps({documents[0].key: {"action": "skip"}})}),
    )
    monkeypatch.setattr(inventory_routes, "write_operation_log", lambda *_args, **_kwargs: None)

    result = asyncio.run(inventory_routes.import_inventory_template(
        request, UploadFile(filename="应付.xlsx", file=BytesIO(content)),
    ))

    assert result["created"] == 1
    assert result["skipped"] == 1
    plans = repository.import_template_documents.call_args.args[0]
    assert plans[0]["decision"] == "skip"
    assert plans[1]["supplier"] == "供应商A"


def test_accounting_template_preview_rejects_unknown_subject():
    content = _workbook(TEMPLATE_HEADERS["accounting"], [
        ["2026-10-01", "应付款增加", "张三", "供应商A", "包装费", "未建科目", 100],
    ])
    repository = Mock(spec=InventoryRepository)
    repository.list_account_subjects.return_value = [{"name": "包装费"}]
    request = SimpleNamespace(app=SimpleNamespace(state=SimpleNamespace(inventory_repository=repository)))

    repository.list_suppliers.return_value = [{"name": "供应商A"}]
    with pytest.raises(HTTPException, match="费用项目名/科目（科目管理）：未建科目") as error:
        asyncio.run(inventory_routes.preview_inventory_template(
            request, UploadFile(filename="应付.xlsx", file=BytesIO(content)),
        ))

    assert error.value.status_code == 400
    repository.preview_template_documents.assert_not_called()


def test_accounting_template_preview_accepts_existing_subject():
    content = _workbook(TEMPLATE_HEADERS["accounting"], [
        ["2026-10-01", "应付款增加", "张三", "供应商A", "包装费", "包装费", 100],
    ])
    repository = Mock(spec=InventoryRepository)
    repository.list_suppliers.return_value = [{"name": "供应商A"}]
    repository.list_account_subjects.return_value = [{"name": "包装费"}]
    repository.preview_template_documents.return_value = {"total": 1, "conflicts": []}
    request = SimpleNamespace(app=SimpleNamespace(state=SimpleNamespace(inventory_repository=repository)))

    result = asyncio.run(inventory_routes.preview_inventory_template(
        request, UploadFile(filename="应付.xlsx", file=BytesIO(content)),
    ))

    assert result["kind"] == "accounting"
    assert result["total"] == 1
    repository.preview_template_documents.assert_called_once()


@pytest.mark.parametrize("kind, values, missing_labels", [
    ("purchase", ["2026-10-01", "进货单", "新供应商", "新仓库", "张三", "进货", "A001", 1, 10],
     ("供应商（供应商管理）：新供应商", "仓库（仓库管理）：新仓库")),
    ("purchase_return", ["2026-10-01", "进货退货单", "新供应商", "新仓库", "张三", "退货", "A001", 1, 10],
     ("供应商（供应商管理）：新供应商", "仓库（仓库管理）：新仓库")),
    ("sale", ["2026-10-01", "批发销售单", "新客户", "新仓库", "张三", "销售", "A001", 1, 10],
     ("一般客户（一般客户管理）：新客户", "仓库（仓库管理）：新仓库")),
    ("sale_return", ["2026-10-01", "批发销售退货单", "新客户", "新仓库", "张三", "退货", "A001", 1, 10],
     ("一般客户（一般客户管理）：新客户", "仓库（仓库管理）：新仓库")),
    ("accounting", ["2026-10-01", "应付款增加", "张三", "新供应商", "应付", "新科目", 100],
     ("供应商（供应商管理）：新供应商", "费用项目名/科目（科目管理）：新科目")),
    ("accounting", ["2026-10-01", "应收款减少", "张三", "新客户", "应收", "新科目", 100],
     ("一般客户（一般客户管理）：新客户", "费用项目名/科目（科目管理）：新科目")),
])
def test_template_preview_lists_all_missing_master_data(kind, values, missing_labels):
    content = _workbook(TEMPLATE_HEADERS[kind], [values])
    repository = Mock(spec=InventoryRepository)
    repository.list_suppliers.return_value = []
    repository.list_warehouses.return_value = []
    repository.list_general_customer_shops.return_value = []
    repository.list_general_customer_units.return_value = []
    repository.list_account_subjects.return_value = []
    request = SimpleNamespace(app=SimpleNamespace(state=SimpleNamespace(inventory_repository=repository)))

    with pytest.raises(HTTPException) as error:
        asyncio.run(inventory_routes.preview_inventory_template(
            request, UploadFile(filename="模板.xlsx", file=BytesIO(content)),
        ))

    assert error.value.status_code == 400
    assert all(label in error.value.detail for label in missing_labels)
    assert "请先在对应管理页面处理后再导入" in error.value.detail
    repository.preview_template_documents.assert_not_called()


def test_template_preview_accepts_general_customer_unit():
    content = _workbook(TEMPLATE_HEADERS["accounting"], [
        ["2026-10-01", "应收款增加", "张三", "客户单位A", "应收", "包装费", 100],
    ])
    repository = Mock(spec=InventoryRepository)
    repository.list_general_customer_shops.return_value = []
    repository.list_general_customer_units.return_value = [{"unit_name": "客户单位A"}]
    repository.list_account_subjects.return_value = [{"name": "包装费"}]
    repository.preview_template_documents.return_value = {"total": 1, "conflicts": []}
    request = SimpleNamespace(app=SimpleNamespace(state=SimpleNamespace(inventory_repository=repository)))

    result = asyncio.run(inventory_routes.preview_inventory_template(
        request, UploadFile(filename="应收.xlsx", file=BytesIO(content)),
    ))

    assert result["total"] == 1
    repository.preview_template_documents.assert_called_once()


def test_accounting_template_import_rechecks_subject_before_writing():
    content = _workbook(TEMPLATE_HEADERS["accounting"], [
        ["2026-10-01", "应收款增加", "张三", "客户A", "包装费", "未建科目", 100],
    ])
    repository = Mock(spec=InventoryRepository)
    repository.list_general_customer_shops.return_value = [{"shop_name": "客户A"}]
    repository.list_general_customer_units.return_value = []
    repository.list_account_subjects.return_value = [{"name": "包装费"}]
    repository.preview_template_documents.return_value = {"conflicts": []}
    request = SimpleNamespace(
        app=SimpleNamespace(state=SimpleNamespace(inventory_repository=repository)),
        form=AsyncMock(return_value={"decisions": "{}"}),
    )

    with pytest.raises(HTTPException, match="费用项目名/科目（科目管理）：未建科目") as error:
        asyncio.run(inventory_routes.import_inventory_template(
            request, UploadFile(filename="应收.xlsx", file=BytesIO(content)),
        ))

    assert error.value.status_code == 400
    repository.import_template_documents.assert_not_called()


def test_accounting_template_repository_rejects_unknown_subject_atomically():
    engine = create_engine("sqlite://")
    event.listen(engine, "connect", lambda connection, _record: connection.create_function(
        "date_trunc", 2, lambda _unit, value: value,
    ))
    INVENTORY_ACCOUNT_SUBJECT_TABLE.create(engine)
    INVENTORY_TABLE.create(engine)
    INVENTORY_DETAIL_TABLE.create(engine)
    SUPPLIER_TABLE.create(engine)
    with engine.begin() as connection:
        connection.execute(SUPPLIER_TABLE.insert().values(id=1, brand="ni", name="供应商A"))
    repository = object.__new__(InventoryRepository)
    repository.engine = engine
    with engine.begin() as connection:
        connection.execute(INVENTORY_ACCOUNT_SUBJECT_TABLE.insert().values(
            id=1, category="费用类", name="包装费",
        ))
    plans = [
        {"document_type": "应付款增加", "supplier": "供应商A", "decision": None, "details": [{"product_name": "包装费", "amount": 100}]},
        {"document_type": "应付款减少", "supplier": "供应商A", "decision": None, "details": [{"product_name": "未建科目", "amount": 50}]},
    ]
    try:
        with pytest.raises(ValueError, match="费用项目名/科目（科目管理）：未建科目"):
            repository.import_template_documents(plans)
        with engine.connect() as connection:
            assert connection.execute(select(INVENTORY_TABLE.c.id)).all() == []
            assert connection.execute(select(INVENTORY_DETAIL_TABLE.c.id)).all() == []
    finally:
        engine.dispose()


def test_template_repository_rejects_missing_customer_and_warehouse_atomically():
    engine = create_engine("sqlite://")
    GENERAL_CUSTOMER_SHOP_TABLE.create(engine)
    GENERAL_CUSTOMER_UNIT_TABLE.create(engine)
    WAREHOUSE_TABLE.create(engine)
    INVENTORY_TABLE.create(engine)
    INVENTORY_DETAIL_TABLE.create(engine)
    repository = object.__new__(InventoryRepository)
    repository.engine = engine
    plans = [{
        "document_type": "批发销售单", "supplier": "新客户", "warehouse": "新仓库",
        "decision": None, "details": [{"product_code": "A001", "quantity": 1, "amount": 100}],
    }]
    try:
        with pytest.raises(ValueError) as error:
            repository.import_template_documents(plans)
        assert "一般客户（一般客户管理）：新客户" in str(error.value)
        assert "仓库（仓库管理）：新仓库" in str(error.value)
        with engine.connect() as connection:
            assert connection.execute(select(INVENTORY_TABLE.c.id)).all() == []
            assert connection.execute(select(INVENTORY_DETAIL_TABLE.c.id)).all() == []
    finally:
        engine.dispose()
